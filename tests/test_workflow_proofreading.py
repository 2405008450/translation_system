from __future__ import annotations

import json
import unittest
from datetime import datetime, timedelta, timezone
from io import BytesIO
from tempfile import TemporaryDirectory
from types import SimpleNamespace
from unittest.mock import patch
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import Session

from app.models import (DocumentAlignmentPair, DocumentAlignmentUnit, FileAssignment, FileRecord, Project, ProjectAssignment, ProjectWorkflowStep,
    ProofreadingBatch, ProofreadingColumnBinding, ProofreadingSegmentBaseline, Segment, SegmentRevision,
    TranslationReviewReport, TranslationReviewReportItem, User)
from app.routers.proofreading import get_file_proofreading_baselines
from app.services.proofreading import generate_batch


class WorkflowTestCase(unittest.IsolatedAsyncioTestCase):
    async def asyncSetUp(self):
        self.storage = TemporaryDirectory()
        self.engine = create_engine("sqlite:///:memory:")
        tables = [User.__table__, Project.__table__, ProjectWorkflowStep.__table__, ProjectAssignment.__table__,
                  FileRecord.__table__, FileAssignment.__table__, Segment.__table__, SegmentRevision.__table__,
                  ProofreadingBatch.__table__, ProofreadingColumnBinding.__table__, ProofreadingSegmentBaseline.__table__,
                  TranslationReviewReport.__table__, TranslationReviewReportItem.__table__,
                  DocumentAlignmentPair.__table__, DocumentAlignmentUnit.__table__]
        defaults = []
        for table in tables:
            for column in table.columns:
                if column.server_default is not None and "::uuid" in str(column.server_default.arg):
                    defaults.append((column, column.server_default))
                    column.server_default = None
        try:
            for table in tables:
                table.create(self.engine)
        finally:
            for column, default in defaults:
                column.server_default = default
        self.db = Session(self.engine)
        self.user = User(username="admin", hashed_password="x", role="admin")
        self.project = Project(name="工作流校对", workflow_template_id="proofread")
        self.db.add_all([self.user, self.project])
        self.db.flush()
        self.db.add(ProjectWorkflowStep(project_id=self.project.id, step_key="proofread", name="校对", step_type="proofread", sort_order=0))
        self.db.commit()

    async def asyncTearDown(self):
        self.db.close()
        self.engine.dispose()
        self.storage.cleanup()


class TranslationWorkflowProofreadingTests(WorkflowTestCase):

    def setup_workflow(self):
        from app.services.workflow_proofreading import prepare_workflow_proofreading

        self.project.workflow_template_id = "translate_review_proofread"
        self.project.source_language = "en-US"
        self.project.target_language = "fr-FR"
        self.proofread_step = self.db.query(ProjectWorkflowStep).filter_by(project_id=self.project.id).one()
        self.proofread_step.sort_order = 2
        self.review_step = ProjectWorkflowStep(project_id=self.project.id, step_key="review", name="审校", step_type="review", sort_order=1)
        self.translation_step = ProjectWorkflowStep(project_id=self.project.id, step_key="translate", name="翻译", step_type="translation", sort_order=0)
        self.db.add_all([self.review_step, self.translation_step])
        self.file = FileRecord(project_id=self.project.id, filename="source.txt", source_language="en-US", target_language="fr-FR")
        self.db.add(self.file)
        self.db.flush()
        self.segments = []
        for index in range(4):
            segment = Segment(
                file_record_id=self.file.id, workflow_step_id=self.proofread_step.id if index < 3 else self.review_step.id,
                sentence_id=f"s{index}", source_text=f"source {index}", display_text=f"source {index}",
                target_text=f"Traduction {index}", source="manual", status="none", display_index=index, sequence_index=index,
            )
            self.db.add(segment)
            self.segments.append(segment)
        self.db.commit()
        batch = prepare_workflow_proofreading(self.db, self.file, self.user)
        self.db.commit()
        return batch

    def assign(self, username, start, end):
        user = User(username=username, hashed_password="x", role="user", translator_type="external")
        self.db.add(user)
        self.db.flush()
        self.db.add(ProjectAssignment(project_id=self.project.id, assignee_id=user.id, status="active"))
        assignment = FileAssignment(project_id=self.project.id, file_record_id=self.file.id,
                                    workflow_step_id=self.proofread_step.id, assignee_id=user.id,
                                    segment_range_start=start, segment_range_end=end, status="active")
        self.db.add(assignment)
        self.db.commit()
        return user, assignment

    async def run_generation(self, batch, user=None, callback=None):
        async def review(groups, **kwargs):
            self.last_groups = groups
            if callback:
                callback()
            return ({group["sid"]: {"reviewed_target_text": "Texte corrigé", "category": "accuracy",
                                      "reason": "统一术语", "confidence": "high"} for group in groups}, "test", "test-model")
        with patch("app.services.proofreading._review_group_batch", side_effect=review), \
             patch("app.services.proofreading._attach_reference_context"), \
             patch("app.services.proofreading.record_translation_metric_event"), \
             patch("app.services.proofreading.sync_file_record_status"):
            await generate_batch(self.db, batch, current_user=user or self.user)

    def test_templates_have_explicit_alignment_and_pure_proofreading_steps(self):
        from app.routers.api import WORKFLOW_TEMPLATE_BY_ID, _create_project_workflow_steps
        self.assertEqual(len(WORKFLOW_TEMPLATE_BY_ID), 5)
        self.assertEqual(WORKFLOW_TEMPLATE_BY_ID["proofread"]["name"], "对齐-校对")
        self.assertEqual([step["step_type"] for step in WORKFLOW_TEMPLATE_BY_ID["proofread"]["steps"]], ["alignment", "proofread"])
        for template_id, expected in [("proofread_only", ["proofread"]), ("proofread", ["alignment", "proofread"])]:
            project = Project(name=template_id, workflow_template_id=template_id)
            self.db.add(project)
            self.db.flush()
            steps = _create_project_workflow_steps(self.db, project, [], template_id)
            self.assertEqual([step.step_type for step in steps], expected)

    def test_complete_schema_still_runs_legacy_stage_data_upgrade(self):
        from unittest.mock import MagicMock
        from app.services import schema_setup
        engine = MagicMock()
        engine.connect.return_value.__enter__.return_value.execute.return_value.scalar.side_effect = [True, False]
        with patch.object(schema_setup, "engine", engine), \
             patch.object(schema_setup, "inspect") as inspector, \
             patch.object(schema_setup, "_collect_missing_schema", return_value=[]), \
             patch.object(schema_setup, "_collect_column_length_migrations", side_effect=lambda _: ([], [])), \
             patch("app.services.review_sync_schema.default_needs_upgrade", return_value=False):
            inspector.return_value.has_table.return_value = True
            schema_setup._ensure_runtime_schema_once()
        statements = [str(call.args[0]) for call in engine.begin.return_value.__enter__.return_value.execute.call_args_list]
        self.assertTrue(any("pg_advisory_xact_lock" in sql for sql in statements))
        self.assertTrue(any("INSERT INTO project_workflow_steps" in sql for sql in statements))
        self.assertTrue(any("UPDATE file_assignments" in sql for sql in statements))
        self.assertFalse(any("CREATE TABLE" in sql for sql in statements))

    def test_pure_proofreading_import_uses_proofreading_step_without_alignment(self):
        from openpyxl import Workbook
        from app.services.proofreading import create_batch_from_workbook
        self.project.workflow_template_id = "proofread_only"
        self.db.commit()
        workbook = Workbook()
        workbook.active.append(["原文", "译文"])
        workbook.active.append(["Hello", "Bonjour"])
        stream = BytesIO()
        workbook.save(stream)
        workbook.close()
        with patch("app.services.proofreading.save_source_file"), \
             patch("app.services.proofreading.attach_project_reference_bases_to_file"):
            batch = create_batch_from_workbook(self.db, project=self.project, current_user=self.user,
                raw_bytes=stream.getvalue(), filename="pure.xlsx", source_language="en-US",
                mappings=[{"sheet_index": 0, "header_row": 1, "source_column": 1,
                           "targets": [{"target_column": 2, "target_language": "fr-FR"}]}])
        self.assertEqual(batch.workflow_stage, "proofreading")
        self.assertEqual(batch.batch_kind, "xlsx_columns")
        segment = self.db.query(Segment).one()
        step = self.db.get(ProjectWorkflowStep, segment.workflow_step_id)
        self.assertEqual(step.step_type, "proofread")

    def test_prepare_is_idempotent_and_preserves_baseline_after_manual_edit(self):
        from app.services.workflow_proofreading import prepare_workflow_proofreading
        batch = self.setup_workflow()
        self.segments[0].target_text = "Modification manuelle"
        self.db.commit()
        same = prepare_workflow_proofreading(self.db, self.file, self.user)
        self.assertEqual(same.id, batch.id)
        self.assertEqual(self.db.query(ProofreadingBatch).count(), 1)
        self.assertEqual(self.db.query(FileRecord).count(), 1)
        self.assertEqual(self.db.query(ProofreadingSegmentBaseline).count(), 3)
        payload = get_file_proofreading_baselines(self.file.id, self.db, self.user)
        self.assertFalse(payload["is_proofreading"])
        self.assertEqual(payload["proofreading_context"]["batch_kind"], "workflow_segments")
        self.assertEqual(payload["items"][0]["original_target_text"], "Traduction 0")
        self.assertNotIn(str(self.segments[3].id), {row["segment_id"] for row in payload["items"]})

    def test_partial_entry_and_reentry_refresh_only_affected_baseline(self):
        from app.services.workflow_proofreading import prepare_workflow_proofreading
        self.setup_workflow()
        self.segments[0].target_text = "Conserver la base"
        self.segments[3].workflow_step_id = self.proofread_step.id
        self.db.flush()
        prepare_workflow_proofreading(self.db, self.file, self.user, refresh_segment_ids={self.segments[3].id})
        self.assertEqual(self.db.query(ProofreadingSegmentBaseline).count(), 4)
        baseline = self.db.query(ProofreadingSegmentBaseline).filter_by(segment_id=self.segments[0].id).one()
        self.assertEqual(baseline.original_target_text, "Traduction 0")
        prepare_workflow_proofreading(self.db, self.file, self.user, refresh_segment_ids={self.segments[0].id})
        self.assertEqual(baseline.original_target_text, "Conserver la base")

    async def test_generation_uses_current_saved_text_and_only_assigned_unconfirmed_segments(self):
        from app.services.workflow_proofreading import eligible_workflow_segments
        batch = self.setup_workflow()
        user, _ = self.assign("proofreader", 1, 2)
        self.segments[0].target_text = "Texte enregistré"
        self.segments[1].status = "confirmed"
        self.db.commit()
        self.assertEqual([segment.id for segment in eligible_workflow_segments(self.db, batch, user)], [self.segments[0].id])
        await self.run_generation(batch, user)
        self.assertEqual(self.last_groups[0]["variants"], ["Texte enregistré"])
        self.assertEqual(self.segments[0].target_text, "Texte corrigé")
        self.assertEqual([segment.target_text for segment in self.segments[1:]], ["Traduction 1", "Traduction 2", "Traduction 3"])
        self.assertEqual(self.db.query(SegmentRevision).count(), 1)
        from app.services.document_alignment.export import export_workflow_proofreading_xlsx
        content, _ = export_workflow_proofreading_xlsx(self.db, batch)
        workbook = load_workbook(BytesIO(content))
        self.assertEqual(workbook.active["G2"].value, "统一术语")
        workbook.close()

    async def test_llm_return_does_not_overwrite_concurrent_edit(self):
        batch = self.setup_workflow()
        def edit():
            self.segments[0].target_text = "Édition pendant la requête"
            self.segments[0].version += 1
            self.db.commit()
        await self.run_generation(batch, callback=edit)
        self.assertEqual(self.segments[0].target_text, "Édition pendant la requête")
        self.assertEqual(self.db.query(SegmentRevision).filter_by(segment_id=self.segments[0].id).count(), 0)

    async def test_revoking_assignment_while_llm_runs_prevents_writeback(self):
        batch = self.setup_workflow()
        user, assignment = self.assign("revoked", 1, 2)
        def revoke():
            assignment.status = "revoked"
            self.db.commit()
        await self.run_generation(batch, user, callback=revoke)
        self.assertEqual(self.segments[0].target_text, "Traduction 0")
        self.assertEqual(self.db.query(SegmentRevision).count(), 0)

    async def test_returning_to_review_while_llm_runs_prevents_writeback(self):
        batch = self.setup_workflow()
        def move_back():
            self.segments[0].workflow_step_id = self.review_step.id
            self.db.commit()
        await self.run_generation(batch, callback=move_back)
        self.assertEqual(self.segments[0].target_text, "Traduction 0")
        payload = get_file_proofreading_baselines(self.file.id, self.db, self.user)
        self.assertNotIn(str(self.segments[0].id), {item["segment_id"] for item in payload["items"]})

    async def test_every_llm_request_runs_without_transaction_or_orm_payload(self):
        from sqlalchemy import inspect
        from app.services.proofreading import _build_prompt
        batch = self.setup_workflow()
        calls = []
        async def review(groups, **kwargs):
            self.assertFalse(self.db.in_transaction())
            for group in groups:
                for baseline, segment in group["segments"]:
                    self.assertIsNone(inspect(baseline, raiseerr=False))
                    self.assertIsNone(inspect(segment, raiseerr=False))
                    self.assertIsNotNone(segment.id)
            _build_prompt(groups, kwargs["source_language"], kwargs["target_language"], kwargs["rules_text"])
            self.assertFalse(self.db.in_transaction())
            calls.append(groups[0]["sid"])
            return ({group["sid"]: {"reviewed_target_text": "Texte corrigé", "category": "accuracy",
                                     "reason": "统一术语", "confidence": "high"} for group in groups}, "test", "model")
        with patch("app.services.proofreading.PROOFREADING_BATCH_SIZE", 1), \
             patch("app.services.proofreading._review_group_batch", side_effect=review), \
             patch("app.services.proofreading._attach_reference_context"), \
             patch("app.services.proofreading.record_translation_metric_event"), \
             patch("app.services.proofreading.sync_file_record_status"):
            await generate_batch(self.db, batch, current_user=self.user)
        self.assertEqual(len(calls), 3)
        self.assertEqual(self.db.query(SegmentRevision).count(), 3)
        self.assertEqual(batch.status, "completed")

    def test_start_endpoint_releases_request_session_before_background_execution(self):
        from fastapi import BackgroundTasks
        from app.routers.proofreading import ProofreadingGenerateRequest, generate_proofreading_batch
        batch = self.setup_workflow()
        batch_id = batch.id
        tasks = BackgroundTasks()
        with patch("app.routers.proofreading.validate_provider_choice"):
            response = generate_proofreading_batch(batch_id, ProofreadingGenerateRequest(), tasks, self.db, self.user)
        self.assertFalse(self.db.in_transaction())
        self.assertEqual(response["batch_id"], str(batch_id))
        self.assertEqual(tasks.tasks[0].args[0], batch_id)

    async def test_database_details_are_hidden_for_new_and_historical_errors(self):
        from app.services.proofreading import DATABASE_RETRY_MESSAGE, serialize_batch
        batch = self.setup_workflow()
        raw_error = "(psycopg.errors.IdleInTransactionSessionTimeout) idle-in-transaction timeout [SQL: SELECT secret FROM segments]"
        with patch("app.services.proofreading._review_group_batch", side_effect=RuntimeError(raw_error)), \
             patch("app.services.proofreading._attach_reference_context"), \
             patch("app.services.proofreading.sync_file_record_status"):
            await generate_batch(self.db, batch, current_user=self.user)
        self.assertEqual(batch.failed_segments, 3)
        self.assertEqual({item.reason for item in self.db.query(TranslationReviewReportItem)}, {DATABASE_RETRY_MESSAGE})
        item = self.db.query(TranslationReviewReportItem).first()
        item.reason = raw_error
        batch.error_message = raw_error
        self.db.commit()
        self.assertEqual(serialize_batch(self.db, batch)["error_message"], DATABASE_RETRY_MESSAGE)
        payload = get_file_proofreading_baselines(self.file.id, self.db, self.user)
        self.assertNotIn("[SQL:", json.dumps(payload, ensure_ascii=False))

    def test_terminal_database_error_finalizes_report_and_keeps_safe_message(self):
        from contextlib import contextmanager
        from unittest.mock import AsyncMock
        from sqlalchemy.exc import SQLAlchemyError
        from app.services.proofreading import DATABASE_RETRY_MESSAGE, run_generate_batch
        batch = self.setup_workflow()
        batch_id, user_id = batch.id, self.user.id
        report = TranslationReviewReport(project_id=self.project.id, proofreading_batch_id=batch.id,
                                         report_mode="proofread_generate", scope="proofreading_batch", status="running")
        self.db.add(report)
        self.db.commit()
        @contextmanager
        def session():
            yield self.db
        with patch("app.services.proofreading.SessionLocal", side_effect=session), \
             patch("app.services.proofreading.generate_batch", new_callable=AsyncMock,
                   side_effect=SQLAlchemyError("connection interrupted [SQL: SELECT ...]")):
            run_generate_batch(batch_id, user_id, "auto", None)
        self.assertEqual(batch.status, "failed")
        self.assertEqual(batch.error_message, DATABASE_RETRY_MESSAGE)
        self.assertEqual(report.status, "failed")
        self.assertIsNotNone(report.finished_at)

    def test_start_is_scoped_busy_is_rejected_and_only_starter_can_cancel(self):
        from fastapi import BackgroundTasks, HTTPException
        from app.routers.proofreading import ProofreadingGenerateRequest, generate_proofreading_batch, cancel_proofreading_batch
        batch = self.setup_workflow()
        user, _ = self.assign("starter", 1, 1)
        other, _ = self.assign("other", 2, 2)
        with patch("app.routers.proofreading.validate_provider_choice"):
            generate_proofreading_batch(batch.id, ProofreadingGenerateRequest(), BackgroundTasks(), self.db, user)
            config = json.loads(batch.config_json)
            self.assertEqual(config["generation"]["segment_ids"], [str(self.segments[0].id)])
            with self.assertRaises(HTTPException) as busy:
                generate_proofreading_batch(batch.id, ProofreadingGenerateRequest(), BackgroundTasks(), self.db, other)
            self.assertEqual(busy.exception.status_code, 409)
        with self.assertRaises(HTTPException) as forbidden:
            cancel_proofreading_batch(batch.id, self.db, other)
        self.assertEqual(forbidden.exception.status_code, 403)
        result = cancel_proofreading_batch(batch.id, self.db, user)
        self.assertEqual(result["status"], "canceled")

    def test_retry_does_not_include_failures_outside_assignment(self):
        from fastapi import BackgroundTasks, HTTPException
        from app.routers.proofreading import ProofreadingGenerateRequest, generate_proofreading_batch
        batch = self.setup_workflow()
        user, _ = self.assign("retry", 1, 1)
        with patch("app.routers.proofreading.get_latest_generation_error_segment_ids", return_value={self.segments[2].id}):
            with self.assertRaises(HTTPException) as forbidden:
                generate_proofreading_batch(batch.id, ProofreadingGenerateRequest(retry_scope="failed_only"), BackgroundTasks(), self.db, user)
        self.assertEqual(forbidden.exception.status_code, 409)

    def test_stale_queued_callback_cannot_claim_restarted_generation(self):
        from contextlib import contextmanager
        from unittest.mock import AsyncMock
        from app.services.proofreading import run_generate_batch
        batch = self.setup_workflow()
        batch.status = "queued"
        config = json.loads(batch.config_json)
        config["generation"] = {"run_id": "new-run"}
        batch.config_json = json.dumps(config)
        self.db.commit()
        @contextmanager
        def session():
            yield self.db
        with patch("app.services.proofreading.SessionLocal", side_effect=session), \
             patch("app.services.proofreading.generate_batch", new_callable=AsyncMock) as generate:
            run_generate_batch(batch.id, self.user.id, "auto", None, run_id="old-run")
        generate.assert_not_awaited()
        self.assertEqual(batch.status, "queued")

    def test_report_and_batch_cleanup_preserve_original_file(self):
        from app.services.document_alignment.export import export_workflow_proofreading_xlsx
        from app.services.file_record_service import delete_proofreading_batch
        batch = self.setup_workflow()
        self.segments[0].source_text = "=SUM(A1)"
        content, filename = export_workflow_proofreading_xlsx(self.db, batch)
        workbook = load_workbook(BytesIO(content))
        self.assertEqual(workbook.active.max_row, 4)
        self.assertEqual(workbook.active["B2"].data_type, "s")
        self.assertEqual(workbook.active["C2"].value, "Traduction 0")
        self.assertTrue(filename.endswith("_校对报告.xlsx"))
        workbook.close()
        with patch("app.services.file_record_service.get_settings", return_value=SimpleNamespace(file_storage_dir=self.storage.name)):
            delete_proofreading_batch(self.db, batch)
        self.assertIsNotNone(self.db.get(FileRecord, self.file.id))
        self.assertEqual(self.db.query(Segment).count(), 4)
        self.assertEqual(self.db.query(ProofreadingSegmentBaseline).count(), 0)

    @staticmethod
    def word_texts(content, accept=True):
        from app.services.document_workspace import NS, _qn
        with ZipFile(BytesIO(content)) as archive:
            root = ET.fromstring(archive.read("word/document.xml"))
        def collect(node):
            if node.tag == _qn("w", "del" if accept else "ins"):
                return ""
            if node.tag in {_qn("w", "t"), _qn("w", "delText")}:
                return node.text or ""
            if node.tag == _qn("w", "tab"):
                return "\t"
            if node.tag == _qn("w", "br"):
                return "\n"
            return "".join(collect(child) for child in node)
        return [collect(paragraph) for paragraph in root.findall(".//w:body//w:p", NS)], root

    def build_revision_export(self, batch, **kwargs):
        from app.services.file_export_queue import build_file_record_exported_file
        return build_file_record_exported_file(
            self.db, self.file, "proofreading_docx_revisions", **kwargs,
        )

    def test_non_word_revision_roundtrip_includes_whole_file_and_confirmed_changes(self):
        batch = self.setup_workflow()
        baselines = {row.segment_id: row for row in self.db.query(ProofreadingSegmentBaseline).all()}
        baselines[self.segments[0].id].original_target_text = "旧文 résumé\n测试\t1"
        self.segments[0].target_text = "新文 résumé\n测试\t2🙂"
        self.segments[0].status = "confirmed"
        baselines[self.segments[1].id].original_target_text = ""
        self.segments[1].target_text = "Inserted text"
        self.segments[2].target_text = ""
        self.segments[3].target_text = "Earlier stage stays clean"
        # 前序修订记录不应进入校对交付文档。
        self.db.add(SegmentRevision(file_record_id=self.file.id, segment_id=self.segments[3].id,
            sentence_id=self.segments[3].sentence_id, before_text="Earlier revision",
            after_text="Earlier stage stays clean", status="pending", source="manual"))
        batch.status = "partial_failed"
        self.db.commit()
        before_count = self.db.query(SegmentRevision).count()
        exported = self.build_revision_export(batch)
        accepted, root = self.word_texts(exported.content)
        rejected, _ = self.word_texts(exported.content, False)
        self.assertEqual(accepted, ["新文 résumé\n测试\t2🙂", "Inserted text", "", "Earlier stage stays clean"])
        self.assertEqual(rejected, ["旧文 résumé\n测试\t1", "", "Traduction 2", "Earlier stage stays clean"])
        from app.services.document_workspace import NS
        self.assertTrue(root.findall(".//w:ins", NS))
        self.assertTrue(root.findall(".//w:del", NS))
        self.assertEqual(exported.filename, "source_校对修订版.docx")
        self.assertEqual(self.db.query(SegmentRevision).count(), before_count)
        self.assertEqual(batch.status, "partial_failed")

    def test_unchanged_word_has_no_revision_marks(self):
        batch = self.setup_workflow()
        exported = self.build_revision_export(batch)
        from app.services.document_workspace import NS
        accepted, root = self.word_texts(exported.content)
        self.assertEqual(accepted, [f"Traduction {index}" for index in range(4)])
        self.assertFalse(root.findall(".//w:ins", NS))
        self.assertFalse(root.findall(".//w:del", NS))

    def test_docx_revision_preserves_layout_table_and_both_translation_versions(self):
        from docx import Document
        from app.services.document_workspace import parse_docx_workspace, NS
        from pathlib import Path
        batch = self.setup_workflow()
        document = Document()
        document.add_heading("Source heading", level=1)
        document.add_paragraph("Source paragraph")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Source cell"
        table.cell(0, 1).text = "Other stage"
        stream = BytesIO()
        document.save(stream)
        raw = stream.getvalue()
        parsed = parse_docx_workspace(raw)["segments"]
        self.assertEqual(len(parsed), 4)
        for segment, source in zip(self.segments, parsed):
            for name in ("sentence_id", "source_text", "display_text", "block_type", "block_index", "row_index", "cell_index"):
                setattr(segment, name, source.get(name))
            metadata = source.get("segment_metadata")
            segment.segment_metadata = json.dumps(metadata) if isinstance(metadata, dict) else metadata
        self.file.filename = "layout.docx"
        self.segments[0].target_text = "Revised heading"
        self.segments[1].target_text = ""
        self.segments[2].target_text = "Corrected cell"
        self.db.commit()
        source_path = Path(self.storage.name) / "layout.docx"
        source_path.write_bytes(raw)
        with patch("app.services.file_record_service.get_file_record_source_path", return_value=source_path), \
             patch("app.services.workflow_proofreading_export.get_file_record_source_path", return_value=source_path):
            exported = self.build_revision_export(batch)
        accepted, root = self.word_texts(exported.content)
        rejected, _ = self.word_texts(exported.content, False)
        self.assertEqual(accepted, ["Revised heading", "", "Corrected cell", "Traduction 3"])
        self.assertEqual(rejected, ["Traduction 0", "Traduction 1", "Traduction 2", "Traduction 3"])
        self.assertEqual(len(root.findall(".//w:tbl", NS)), 1)
        self.assertEqual(root.find(".//w:p/w:pPr/w:pStyle", NS).get(f"{{{NS['w']}}}val"), "Heading1")
        self.assertEqual(source_path.read_bytes(), raw)

    def test_revision_snapshot_survives_manual_edit_and_releases_transaction_before_render(self):
        from app.services import workflow_proofreading_export as exporter
        batch = self.setup_workflow()
        self.segments[0].target_text = "Snapshot text"
        self.db.commit()
        render = exporter.render_workflow_proofreading_docx
        def check(snapshot):
            self.assertFalse(self.db.in_transaction())
            self.assertIsInstance(snapshot.segments[0], dict)
            # 渲染使用快照，不读取后台执行过程中发生的其他修改。
            with Session(self.engine) as other:
                other.get(Segment, snapshot.segments[0]["id"]).target_text = "Later saved edit"
                other.commit()
            return render(snapshot)
        with patch.object(exporter, "render_workflow_proofreading_docx", side_effect=check):
            exported = self.build_revision_export(batch, release_transaction_before_render=True)
        self.assertEqual(self.word_texts(exported.content)[0][0], "Snapshot text")
        self.assertFalse(self.db.in_transaction())

    def test_revision_export_rejects_missing_baseline_and_duplicate_mapping(self):
        batch = self.setup_workflow()
        baseline = self.db.query(ProofreadingSegmentBaseline).filter_by(segment_id=self.segments[0].id).one()
        self.db.delete(baseline)
        self.db.commit()
        with self.assertRaisesRegex(ValueError, "缺少比较基线"):
            self.build_revision_export(batch)
        self.segments[0].workflow_step_id = self.review_step.id
        self.segments[1].sentence_id = self.segments[0].sentence_id
        self.db.commit()
        with self.assertRaisesRegex(ValueError, "标识缺失或重复"):
            self.build_revision_export(batch)

    def test_readiness_busy_failed_cancelled_and_missing_word_source(self):
        from app.services.document_alignment.export import build_export_readiness
        batch = self.setup_workflow()
        for status in ("ready", "completed", "partial_failed", "failed", "canceled"):
            batch.status = status
            self.assertIn("proofreading_docx_revisions", build_export_readiness(self.db, batch)["available_formats"])
        for status in ("queued", "running", "canceling"):
            batch.status = status
            readiness = build_export_readiness(self.db, batch)
            self.assertEqual(readiness["available_formats"], [])
            self.assertIn("proofreading_docx_revisions", readiness["unavailable_formats"])
            with self.assertRaisesRegex(ValueError, "再导出"):
                self.build_revision_export(batch)
        batch.status = "ready"
        self.file.filename = "missing.docx"
        with patch("app.services.workflow_proofreading_export.get_file_record_source_path", return_value=None):
            readiness = build_export_readiness(self.db, batch)
            self.assertIn("proofreading_audit_xlsx", readiness["available_formats"])
            self.assertNotIn("proofreading_docx_revisions", readiness["available_formats"])
            self.assertIn("源文件缺失", readiness["unavailable_formats"]["proofreading_docx_revisions"])
            with self.assertRaisesRegex(ValueError, "源文件缺失"):
                self.build_revision_export(batch)

    def test_revision_export_uses_refreshed_baseline_after_return_and_reentry(self):
        from app.services.workflow_proofreading import prepare_workflow_proofreading
        batch = self.setup_workflow()
        self.segments[0].target_text = "First proofread pass"
        self.segments[0].workflow_step_id = self.review_step.id
        self.db.commit()
        exported = self.build_revision_export(batch)
        self.assertEqual(self.word_texts(exported.content, False)[0][0], "First proofread pass")
        self.segments[0].target_text = "New review baseline"
        self.segments[0].workflow_step_id = self.proofread_step.id
        prepare_workflow_proofreading(self.db, self.file, self.user, refresh_segment_ids={self.segments[0].id})
        self.segments[0].target_text = "Second proofreading result"
        self.db.commit()
        exported = self.build_revision_export(batch)
        self.assertEqual(self.word_texts(exported.content, False)[0][0], "New review baseline")
        self.assertEqual(self.word_texts(exported.content)[0][0], "Second proofreading result")

    def test_excel_report_releases_transaction_and_preserves_columns_and_order(self):
        from app.services.file_export_queue import build_file_record_exported_file
        batch = self.setup_workflow()
        self.segments[0].target_text = "Saved edit"
        self.segments[0].status = "confirmed"
        self.db.commit()
        exported = build_file_record_exported_file(
            self.db, self.file, "proofreading_audit_xlsx", release_transaction_before_render=True,
        )
        self.assertFalse(self.db.in_transaction())
        workbook = load_workbook(BytesIO(exported.content))
        rows = list(workbook.active.values)
        self.assertEqual(rows[0], ("序号", "原文", "进入校对时译文", "当前译文", "是否修改", "确认状态", "修改建议或失败原因"))
        self.assertEqual(len(rows), 4)
        self.assertEqual(rows[1][2:6], ("Traduction 0", "Saved edit", "是", "已确认"))
        self.assertEqual([row[1] for row in rows[1:]], [f"source {index}" for index in range(3)])
        workbook.close()

    def test_revision_export_endpoint_checks_access_busy_state_and_warning_ack(self):
        from fastapi import HTTPException
        from app.routers.proofreading import (
            ProofreadingExportTaskRequest, create_proofreading_export_task, get_proofreading_export_readiness,
        )
        batch = self.setup_workflow()
        outsider = User(username="unassigned", hashed_password="x", role="user", translator_type="external")
        self.db.add(outsider)
        self.db.commit()
        payload = ProofreadingExportTaskRequest(format="proofreading_docx_revisions", acknowledge_warnings=True)
        with patch("app.routers.proofreading.queue_file_export", return_value={"task_id": "test"}) as queue:
            with self.assertRaises(HTTPException) as error:
                get_proofreading_export_readiness(batch.id, self.db, outsider)
            self.assertEqual(error.exception.status_code, 404)
            with self.assertRaises(HTTPException) as error:
                create_proofreading_export_task(batch.id, payload, self.db, outsider)
            self.assertEqual(error.exception.status_code, 404)
            queue.assert_not_called()
            assigned, _ = self.assign("assigned-exporter", 1, 2)
            batch.status = "running"
            with self.assertRaises(HTTPException) as error:
                create_proofreading_export_task(batch.id, payload, self.db, assigned)
            self.assertEqual(error.exception.status_code, 409)
            batch.status = "failed"
            with self.assertRaises(HTTPException) as error:
                create_proofreading_export_task(batch.id,
                    ProofreadingExportTaskRequest(format="proofreading_docx_revisions"), self.db, assigned)
            self.assertEqual(error.exception.status_code, 409)
            response = create_proofreading_export_task(batch.id, payload, self.db, assigned)
            self.assertEqual(response.status_code, 202)
            self.assertEqual(queue.call_args.kwargs["export_type"], "proofreading_docx_revisions")
            self.assertEqual(queue.call_args.kwargs["file_record_id"], self.file.id)

    def test_docx_hyperlink_changes_cannot_silently_lose_revision_marks(self):
        from docx import Document
        from docx.oxml import OxmlElement
        from app.services.document_workspace import parse_docx_workspace, _qn
        from app.services.task_file_service import export_translated_task_file
        document = Document()
        paragraph = document.add_paragraph()
        hyperlink = OxmlElement("w:hyperlink")
        hyperlink.set(_qn("w", "anchor"), "bookmark")
        run = OxmlElement("w:r")
        text = OxmlElement("w:t")
        text.text = "Source hyperlink text"
        run.append(text)
        hyperlink.append(run)
        paragraph._p.append(hyperlink)
        stream = BytesIO()
        document.save(stream)
        raw = stream.getvalue()
        segments = parse_docx_workspace(raw)["segments"]
        self.assertEqual(len(segments), 1)
        segments[0]["target_text"] = "Corrected translation"
        revisions = [{"id": "proofreading-test", "sentence_id": segments[0]["sentence_id"],
                      "before_text": "Review translation", "after_text": "Corrected translation",
                      "status": "pending", "source": "manual"}]
        with self.assertRaisesRegex(ValueError, "无法安全写入真实校对修订"):
            export_translated_task_file(raw, "hyperlink.docx", segments, revisions=revisions,
                include_revision_marks=True, require_revision_marks=True)
        # 原有普通翻译导出继续保留既有降级行为。
        self.assertTrue(export_translated_task_file(raw, "hyperlink.docx", segments,
            revisions=revisions, include_revision_marks=True).content)

    def test_excel_and_ppt_sources_export_ordered_revision_word_without_source_file(self):
        batch = self.setup_workflow()
        self.segments[0].target_text = "Corrected translation"
        self.db.commit()
        for filename in ("sheet.xlsx", "slides.pptx"):
            self.file.filename = filename
            self.db.commit()
            with patch("app.services.workflow_proofreading_export.get_file_record_source_path", return_value=None):
                exported = self.build_revision_export(batch)
            self.assertEqual(self.word_texts(exported.content)[0],
                ["Corrected translation", "Traduction 1", "Traduction 2", "Traduction 3"])
            self.assertEqual(self.word_texts(exported.content, False)[0],
                [f"Traduction {index}" for index in range(4)])

    def add_suggestion(self, batch, segment, text, reason="统一术语，并修正语法。", *, category="语言与格式", created_at=None, status="applied"):
        created_at = created_at or datetime.now(timezone.utc).replace(tzinfo=None) + timedelta(seconds=1)
        report = TranslationReviewReport(proofreading_batch_id=batch.id, file_record_id=self.file.id,
            project_id=self.project.id, created_at=created_at, status="completed")
        self.db.add(report)
        self.db.flush()
        item = TranslationReviewReportItem(report_id=report.id, file_record_id=self.file.id,
            segment_id=segment.id, sentence_id=segment.sentence_id, category_key=category,
            suggested_target_text=text, reason=reason, applied=status == "applied", status=status,
            created_at=created_at)
        self.db.add(item)
        self.db.flush()
        return item

    def word_comments(self, content):
        from lxml import etree
        from app.services.docx_comments import W, COMMENTS_REL, DOCUMENT_RELS, COMMENTS_TYPE
        import posixpath
        with ZipFile(BytesIO(content)) as archive:
            self.assertIsNone(archive.testzip())
            rels = etree.fromstring(archive.read(DOCUMENT_RELS))
            comment_rels = [node for node in rels if node.get("Type") == COMMENTS_REL]
            if not comment_rels:
                return {}
            self.assertEqual(len(comment_rels), 1)
            target = comment_rels[0].get("Target")
            part = posixpath.normpath(posixpath.join("word", target))
            types = etree.fromstring(archive.read("[Content_Types].xml"))
            self.assertTrue(any(node.get("PartName") == "/" + part and node.get("ContentType") == COMMENTS_TYPE for node in types))
            comments = etree.fromstring(archive.read(part))
            root = etree.fromstring(archive.read("word/document.xml"))
        def collect(node, accept=True):
            if node.tag == W + ("del" if accept else "ins"):
                return ""
            if node.tag in {W + "t", W + "delText"}:
                return node.text or ""
            if node.tag == W + "br":
                return "\n"
            return "".join(collect(child, accept) for child in node)
        result = {}
        for comment in comments:
            comment_id = comment.get(W + "id")
            start = root.findall(f".//{W}commentRangeStart[@{W}id='{comment_id}']")
            end = root.findall(f".//{W}commentRangeEnd[@{W}id='{comment_id}']")
            reference = root.findall(f".//{W}commentReference[@{W}id='{comment_id}']")
            self.assertEqual((len(start), len(end), len(reference)), (1, 1, 1))
            paragraph = start[0].getparent()
            self.assertIs(paragraph, end[0].getparent())
            self.assertIs(paragraph, reference[0].getparent().getparent())
            self.assertFalse(any(node.tag in {W + "ins", W + "del"} for node in start[0].iterancestors()))
            nodes = list(paragraph)[paragraph.index(start[0]) + 1:paragraph.index(end[0])]
            result[comment_id] = {
                "text": collect(comment), "author": comment.get(W + "author"),
                "accepted": "".join(collect(node) for node in nodes),
                "rejected": "".join(collect(node, False) for node in nodes),
            }
        return result

    async def test_llm_suggestions_become_whole_sentence_comments_in_ordered_word(self):
        batch = self.setup_workflow()
        await self.run_generation(batch)
        exported = self.build_revision_export(batch, release_transaction_before_render=True)
        self.assertFalse(self.db.in_transaction())
        comments = self.word_comments(exported.content)
        self.assertEqual(len(comments), 3)
        self.assertEqual({item["text"] for item in comments.values()}, {"accuracy：统一术语"})
        self.assertEqual({item["author"] for item in comments.values()}, {"校对工作流"})
        self.assertEqual({item["accepted"] for item in comments.values()}, {"Texte corrigé"})
        self.assertEqual({item["rejected"] for item in comments.values()}, {f"Traduction {index}" for index in range(3)})
        self.assertEqual(self.word_texts(exported.content)[0][-1], "Traduction 3")

    def test_comments_keep_latest_success_through_failure_cancel_and_retry(self):
        batch = self.setup_workflow()
        self.segments[0].target_text = "Latest successful text"
        self.segments[1].target_text = "Still valid earlier text"
        moment = datetime.now(timezone.utc).replace(tzinfo=None) + timedelta(seconds=1)
        self.add_suggestion(batch, self.segments[0], "Older text", "过期理由", created_at=moment)
        self.add_suggestion(batch, self.segments[0], "Latest successful text", "最新成功理由", created_at=moment + timedelta(seconds=1))
        self.add_suggestion(batch, self.segments[1], "Still valid earlier text", "另一句仍有效的理由", created_at=moment)
        self.add_suggestion(batch, self.segments[0], "", "数据库异常不进入批注", category="generation_error",
            created_at=moment + timedelta(seconds=2), status="open")
        self.db.commit()
        for state in ("failed", "canceled", "partial_failed", "completed"):
            batch.status = state
            self.db.commit()
            comments = self.word_comments(self.build_revision_export(batch).content)
            self.assertEqual({comment["text"] for comment in comments.values()},
                {"语言与格式：最新成功理由", "语言与格式：另一句仍有效的理由"})

    def test_stale_manual_unchanged_empty_and_resolved_reasons_do_not_create_comments(self):
        batch = self.setup_workflow()
        self.segments[0].target_text = "Saved LLM result"
        self.add_suggestion(batch, self.segments[0], "Saved LLM result")
        self.segments[0].target_text = "Later manual edit"
        self.add_suggestion(batch, self.segments[1], "Traduction 1", "无需修改", category="无需修改")
        self.segments[2].target_text = "Changed without valid reason"
        self.add_suggestion(batch, self.segments[2], self.segments[2].target_text, "", status="applied")
        self.add_suggestion(batch, self.segments[2], self.segments[2].target_text, "已解决的建议", status="resolved")
        self.db.commit()
        exported = self.build_revision_export(batch)
        self.assertEqual(self.word_comments(exported.content), {})
        from app.services.document_workspace import NS
        self.assertTrue(self.word_texts(exported.content)[1].findall(".//w:ins", NS))

    def test_comments_exclude_previous_baseline_after_reentry(self):
        from app.services.workflow_proofreading import prepare_workflow_proofreading
        batch = self.setup_workflow()
        old_time = datetime.now(timezone.utc).replace(tzinfo=None) - timedelta(seconds=1)
        self.add_suggestion(batch, self.segments[0], "Final text", "旧基线的理由", created_at=old_time)
        self.segments[0].target_text = "Reentered baseline"
        prepare_workflow_proofreading(self.db, self.file, self.user, refresh_segment_ids={self.segments[0].id})
        self.segments[0].target_text = "Final text"
        self.db.commit()
        self.assertEqual(self.word_comments(self.build_revision_export(batch).content), {})
        self.add_suggestion(batch, self.segments[0], "Final text", "重入后的新理由")
        self.db.commit()
        comments = self.word_comments(self.build_revision_export(batch).content)
        self.assertEqual(list(comments.values())[0]["text"], "语言与格式：重入后的新理由")
        self.assertEqual(list(comments.values())[0]["rejected"], "Reentered baseline")

    def test_comments_support_full_insertion_deletion_and_confirmed_segments(self):
        batch = self.setup_workflow()
        baseline = self.db.query(ProofreadingSegmentBaseline).filter_by(segment_id=self.segments[0].id).one()
        baseline.original_target_text = ""
        self.segments[0].target_text = "Inserted whole sentence"
        self.segments[0].status = "confirmed"
        self.segments[1].target_text = ""
        self.add_suggestion(batch, self.segments[0], "Inserted whole sentence", "补译整句")
        self.add_suggestion(batch, self.segments[1], "", "删除重复句")
        self.db.commit()
        comments = list(self.word_comments(self.build_revision_export(batch).content).values())
        self.assertEqual([(row["accepted"], row["rejected"]) for row in comments],
            [("Inserted whole sentence", ""), ("", "Traduction 1")])

    def test_docx_whole_sentence_comments_preserve_existing_comments_table_and_revision_roundtrip(self):
        from docx import Document
        from lxml import etree
        from pathlib import Path
        from app.services.document_workspace import parse_docx_workspace
        from app.services.docx_comments import W, COMMENTS_REL, COMMENTS_TYPE, DOCUMENT_RELS, RELS, TYPES
        batch = self.setup_workflow()
        document = Document()
        document.add_paragraph("First source sentence. Second source sentence.")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Source table text"
        table.cell(0, 1).text = "Other stage text"
        stream = BytesIO()
        document.save(stream)
        # 手工建立含原有批注、非默认批注路径的母版，不依赖待验证的批注写入器。
        raw_stream = BytesIO()
        with ZipFile(BytesIO(stream.getvalue())) as archive, ZipFile(raw_stream, "w") as output:
            body = etree.fromstring(archive.read("word/document.xml"))
            paragraph = body.find(".//" + W + "body/" + W + "p")
            paragraph.insert(0, etree.Element(W + "commentRangeStart", {W + "id": "41"}))
            paragraph.append(etree.Element(W + "commentRangeEnd", {W + "id": "41"}))
            reference = etree.SubElement(paragraph, W + "r")
            etree.SubElement(reference, W + "commentReference", {W + "id": "41"})
            old_comments = etree.Element(W + "comments", nsmap={"w": W[1:-1]})
            old = etree.SubElement(old_comments, W + "comment", {W + "id": "41", W + "author": "原有审阅人"})
            etree.SubElement(etree.SubElement(etree.SubElement(old, W + "p"), W + "r"), W + "t").text = "母版已有批注"
            rels = etree.fromstring(archive.read(DOCUMENT_RELS))
            etree.SubElement(rels, f"{{{RELS}}}Relationship", Id="rIdExistingComments", Type=COMMENTS_REL, Target="reviewer-notes.xml")
            types = etree.fromstring(archive.read("[Content_Types].xml"))
            etree.SubElement(types, f"{{{TYPES}}}Override", PartName="/word/reviewer-notes.xml", ContentType=COMMENTS_TYPE)
            updates = {"word/document.xml": body, DOCUMENT_RELS: rels, "[Content_Types].xml": types}
            for info in archive.infolist():
                output.writestr(info, etree.tostring(updates[info.filename]) if info.filename in updates else archive.read(info.filename))
            output.writestr("word/reviewer-notes.xml", etree.tostring(old_comments))
        raw = raw_stream.getvalue()
        options = {"include_comments": False}
        parsed = parse_docx_workspace(raw, document_parse_options=options)["segments"]
        self.assertEqual(len(parsed), 4)
        for segment, source in zip(self.segments, parsed):
            for name in ("sentence_id", "source_text", "display_text", "block_type", "block_index", "row_index", "cell_index"):
                setattr(segment, name, source.get(name))
            segment.segment_metadata = json.dumps(source.get("segment_metadata") or {})
        self.file.filename = "source-with-comments.docx"
        self.file.document_parse_options = json.dumps(options)
        for index, segment in enumerate(self.segments[:3]):
            segment.target_text = ["Revised first sentence.", "Second revised sentence.", "Corrected table text"][index]
            self.add_suggestion(batch, segment, segment.target_text, f"第{index + 1}句的整句修改建议")
        self.db.commit()
        path = Path(self.storage.name) / self.file.filename
        path.write_bytes(raw)
        with patch("app.services.file_record_service.get_file_record_source_path", return_value=path), \
             patch("app.services.workflow_proofreading_export.get_file_record_source_path", return_value=path):
            exported = self.build_revision_export(batch, release_transaction_before_render=True)
        self.assertFalse(self.db.in_transaction())
        comments = self.word_comments(exported.content)
        self.assertEqual(set(comments), {"41", "42", "43", "44"})
        self.assertEqual(comments["41"]["text"], "母版已有批注")
        self.assertEqual(comments["41"]["author"], "原有审阅人")
        self.assertEqual([comments[str(index)]["accepted"].strip() for index in (42, 43, 44)],
            ["Revised first sentence.", "Second revised sentence.", "Corrected table text"])
        self.assertEqual([comments[str(index)]["rejected"].strip() for index in (42, 43, 44)],
            ["Traduction 0", "Traduction 1", "Traduction 2"])
        self.assertEqual(comments["44"]["text"], "语言与格式：第3句的整句修改建议")
        self.assertEqual(self.word_texts(exported.content)[0],
            ["Revised first sentence. Second revised sentence.", "Corrected table text", "Traduction 3"])
        self.assertEqual(path.read_bytes(), raw)

    def test_standalone_excel_word_comments_follow_language_columns_and_saved_text(self):
        from openpyxl import Workbook
        from app.services.proofreading import create_batch_from_workbook
        from app.services.document_alignment.export import export_ordered_bilingual_docx, build_export_readiness
        self.project.workflow_template_id = "proofread_only"
        self.db.commit()
        workbook = Workbook()
        workbook.active.append(["原文", "法语", "英语"])
        workbook.active.append(["重复原文", "Avant FR", "Before EN"])
        workbook.active.append(["重复原文", "Autre FR", "Other EN"])
        stream = BytesIO()
        workbook.save(stream)
        workbook.close()
        with patch("app.services.proofreading.save_source_file"), \
             patch("app.services.proofreading.attach_project_reference_bases_to_file"):
            batch = create_batch_from_workbook(self.db, project=self.project, current_user=self.user,
                raw_bytes=stream.getvalue(), filename="pure.xlsx", source_language="zh-CN",
                mappings=[{"sheet_index": 0, "header_row": 1, "source_column": 1,
                    "targets": [{"target_column": 2, "target_language": "fr-FR"},
                                {"target_column": 3, "target_language": "en-US"}]}])
        bindings = self.db.query(ProofreadingColumnBinding).filter_by(batch_id=batch.id).all()
        for binding in bindings:
            self.file = self.db.get(FileRecord, binding.file_record_id)
            segments = self.db.query(Segment).filter_by(file_record_id=self.file.id).order_by(Segment.sequence_index).all()
            segments[0].target_text = f"Corrected {binding.target_language}"
            self.add_suggestion(batch, segments[0], segments[0].target_text, f"统一{binding.target_language}表达")
            self.add_suggestion(batch, segments[1], "LLM previous text", "人工改写后不再适用")
            segments[1].target_text = "Manual final text"
        self.db.commit()
        readiness = build_export_readiness(self.db, batch)
        self.assertIn("proofreading_docx_ordered", readiness["available_formats"])
        content, _ = export_ordered_bilingual_docx(self.db, batch, release_transaction_before_render=True)
        self.assertFalse(self.db.in_transaction())
        comments = list(self.word_comments(content).values())
        self.assertEqual([row["accepted"] for row in comments], ["Corrected fr-FR", "Corrected en-US"])
        self.assertEqual([row["text"] for row in comments], ["语言与格式：统一fr-FR表达", "语言与格式：统一en-US表达"])
        self.assertEqual({row["author"] for row in comments}, {"校对工作流"})

    def setup_standalone_document_pair(self):
        from docx import Document
        from app.services.document_alignment.parser import parse_side
        from app.services.document_alignment.segments import ensure_document_pair_segments_complete
        document = Document()
        document.add_paragraph("First sentence. Second sentence.")
        document.add_table(rows=1, cols=1).cell(0, 0).text = "Table text."
        stream = BytesIO()
        document.save(stream)
        raw = stream.getvalue()
        units = parse_side(raw, "target.docx", "sentence")
        self.assertEqual(len(units), 3)
        batch = ProofreadingBatch(project_id=self.project.id, created_by_id=self.user.id,
            filename="source.docx", file_hash="test", source_language="en-US", target_language="fr-FR",
            batch_kind="document_pair", alignment_status="confirmed", workflow_stage="proofreading",
            status="completed", config_json=json.dumps({"target_filename": "target.docx", "granularity": "sentence"}))
        self.file = FileRecord(project_id=self.project.id, filename="source.docx", source_language="en-US",
            target_language="fr-FR", document_parse_mode="full", document_parse_options="{}")
        self.db.add_all([batch, self.file])
        self.db.flush()
        self.db.add(ProofreadingColumnBinding(batch_id=batch.id, file_record_id=self.file.id,
            sheet_index=0, sheet_name="双文档对齐", header_row=1, source_column=1, target_column=2,
            output_column=3, target_language="fr-FR"))
        pairs = []
        for unit in units:
            pair = DocumentAlignmentPair(batch_id=batch.id, pair_order=unit.index,
                src_indices=json.dumps([unit.index]), tgt_indices=json.dumps([unit.index]),
                source_text=unit.text, target_text=unit.text, block_type=unit.block_type,
                block_index=unit.block_index, row_index=unit.row_index, cell_index=unit.cell_index,
                confidence=1, confidence_level="high", method="manual")
            self.db.add(pair)
            pairs.append(pair)
        self.db.flush()
        ensure_document_pair_segments_complete(self.db, batch)
        self.segments = self.db.query(Segment).filter_by(file_record_id=self.file.id).order_by(Segment.sequence_index).all()
        self.assertEqual(len(self.segments), 3)
        for index, segment in enumerate(self.segments):
            segment.target_text = ["First corrected sentence.", "Second corrected sentence.", "Corrected table text."][index]
            self.add_suggestion(batch, segment, segment.target_text, f"第{index + 1}句中文理由")
        self.db.commit()
        return batch, raw

    def test_standalone_target_word_comments_cover_sentences_table_and_revision_roundtrip(self):
        from pathlib import Path
        from app.services.document_alignment.export import export_target_docx_with_revisions
        batch, raw = self.setup_standalone_document_pair()
        path = Path(self.storage.name) / "target.docx"
        path.write_bytes(raw)
        with patch("app.services.document_alignment.export.target_cache_path", return_value=path):
            content, _ = export_target_docx_with_revisions(self.db, batch, release_transaction_before_render=True)
        self.assertFalse(self.db.in_transaction())
        comments = list(self.word_comments(content).values())
        self.assertEqual([row["accepted"].strip() for row in comments],
            ["First corrected sentence.", "Second corrected sentence.", "Corrected table text."])
        self.assertEqual([row["rejected"].strip() for row in comments], ["First sentence.", "Second sentence.", "Table text."])
        self.assertEqual([row["text"] for row in comments], [f"语言与格式：第{index}句中文理由" for index in (1, 2, 3)])
        self.assertEqual(self.word_texts(content)[0], ["First corrected sentence. Second corrected sentence.", "Corrected table text."])
        self.assertEqual(self.word_texts(content, accept=False)[0], ["First sentence. Second sentence.", "Table text."])
        self.assertEqual(path.read_bytes(), raw)

    def test_standalone_layout_and_fallback_keep_comments_without_extra_revisions(self):
        from app.services.document_alignment.export import export_layout_bilingual_docx
        from app.services.document_workspace import NS
        batch, raw = self.setup_standalone_document_pair()
        with patch("app.services.file_record_service.load_file_record_source", return_value=raw):
            content, _, fallback = export_layout_bilingual_docx(self.db, batch, release_transaction_before_render=True)
        self.assertFalse(self.db.in_transaction())
        self.assertFalse(fallback)
        self.assertEqual(len(self.word_comments(content)), 3)
        root = self.word_texts(content)[1]
        self.assertFalse(root.findall(".//w:ins", NS))
        self.assertFalse(root.findall(".//w:del", NS))
        with patch("app.services.file_record_service.load_file_record_source", return_value=raw), \
             patch("app.services.task_file_service.export_bilingual_task_docx_with_layout", side_effect=ValueError("无法定位")):
            content, _, fallback = export_layout_bilingual_docx(self.db, self.db.get(ProofreadingBatch, batch.id))
        self.assertTrue(fallback)
        self.assertEqual(len(self.word_comments(content)), 3)

    def test_standalone_word_queue_keeps_snapshot_and_releases_transaction_before_render(self):
        from pathlib import Path
        from app.services import document_exporter
        from app.services.file_export_queue import build_file_record_exported_file
        batch, raw = self.setup_standalone_document_pair()
        file_id, batch_id = self.file.id, batch.id
        path = Path(self.storage.name) / "target.docx"
        path.write_bytes(raw)
        renderer = document_exporter.export_translated_docx
        def render_after_release(*args, **kwargs):
            self.assertFalse(self.db.in_transaction())
            return renderer(*args, **kwargs)
        with patch("app.services.document_alignment.export.target_cache_path", return_value=path), \
             patch("app.services.document_exporter.export_translated_docx", side_effect=render_after_release):
            result = build_file_record_exported_file(self.db, self.file, "proofreading_docx_target_revisions",
                release_transaction_before_render=True)
        self.assertEqual(len(self.word_comments(result.content)), 3)
        file_record = self.db.get(FileRecord, file_id)
        result = build_file_record_exported_file(self.db, file_record, "proofreading_docx_ordered",
            release_transaction_before_render=True)
        self.assertFalse(self.db.in_transaction())
        self.assertEqual(len(self.word_comments(result.content)), 3)
        self.assertEqual(self.db.query(TranslationReviewReportItem).count(), 3)
        self.assertEqual(self.db.query(ProofreadingSegmentBaseline).filter_by(batch_id=batch_id).count(), 3)

    def test_standalone_target_word_whole_insert_delete_and_existing_comment_ids(self):
        from pathlib import Path
        from docx import Document
        from app.services.document_alignment.export import export_target_docx_with_revisions
        batch, raw = self.setup_standalone_document_pair()
        first_base = self.db.query(ProofreadingSegmentBaseline).filter_by(segment_id=self.segments[0].id).one()
        first_base.original_target_text = ""
        self.segments[1].target_text = ""
        self.add_suggestion(batch, self.segments[1], "", "删除整句重复内容")
        self.db.commit()
        document = Document(BytesIO(raw))
        document.add_comment(document.paragraphs[0].runs[0], text="母版原有批注", author="原审阅人")
        stream = BytesIO()
        document.save(stream)
        path = Path(self.storage.name) / "target.docx"
        path.write_bytes(stream.getvalue())
        with patch("app.services.document_alignment.export.target_cache_path", return_value=path):
            content, _ = export_target_docx_with_revisions(self.db, batch)
        comments = self.word_comments(content)
        self.assertEqual(set(comments), {"0", "1", "2", "3"})
        self.assertEqual(comments["0"]["text"], "母版原有批注")
        self.assertEqual(comments["0"]["author"], "原审阅人")
        self.assertEqual((comments["1"]["accepted"].strip(), comments["1"]["rejected"].strip()),
            ("First corrected sentence.", ""))
        self.assertEqual((comments["2"]["accepted"].strip(), comments["2"]["rejected"].strip()),
            ("", "Second sentence."))
        self.assertEqual(comments["2"]["text"], "语言与格式：删除整句重复内容")

    def test_standalone_target_word_merged_pair_comment_maps_to_each_changed_sentence(self):
        from app.services.document_alignment.export import _build_target_revision_payload
        batch, raw = self.setup_standalone_document_pair()
        pairs = self.db.query(DocumentAlignmentPair).filter_by(batch_id=batch.id).order_by(DocumentAlignmentPair.pair_order).all()
        pairs[0].tgt_indices = '["0", "1"]'
        pairs[0].target_text = "First sentence.\nSecond sentence."
        pairs[1].tgt_indices = "[]"
        pairs[1].target_text = ""
        self.segments[0].target_text = "First corrected sentence.\nSecond corrected sentence."
        self.segments[1].target_text = ""
        baseline_by_id = {item.segment_id: item for item in self.db.query(ProofreadingSegmentBaseline).all()}
        baseline_by_id[self.segments[0].id].original_target_text = pairs[0].target_text
        baseline_by_id[self.segments[1].id].original_target_text = ""
        self.add_suggestion(batch, self.segments[0], self.segments[0].target_text, "合并对齐对的理由")
        self.db.commit()
        _, revisions = _build_target_revision_payload(self.db, batch, raw, "target.docx")
        self.assertEqual([row["comment_text"] for row in revisions],
            ["语言与格式：合并对齐对的理由", "语言与格式：合并对齐对的理由", "语言与格式：第3句中文理由"])
        with patch("app.services.document_alignment.export._apply_unit_replacements_to_block", side_effect=lambda parts, *args, **kwargs: list(parts)):
            with self.assertRaisesRegex(ValueError, "无法安全定位"):
                _build_target_revision_payload(self.db, batch, raw, "target.docx")

    def test_standalone_alignment_stage_and_refreshed_baseline_exclude_old_comments(self):
        from app.services.document_alignment.export import export_ordered_bilingual_docx
        batch, _ = self.setup_standalone_document_pair()
        batch.workflow_stage = "alignment"
        self.db.commit()
        self.assertEqual(self.word_comments(export_ordered_bilingual_docx(self.db, batch)[0]), {})
        batch.workflow_stage = "proofreading"
        baseline = self.db.query(ProofreadingSegmentBaseline).filter_by(segment_id=self.segments[0].id).one()
        baseline.created_at = datetime.now(timezone.utc).replace(tzinfo=None) + timedelta(seconds=5)
        self.segments[1].target_text = "Manual text invalidates old suggestion"
        self.add_suggestion(batch, self.segments[2], "", "异常信息", category="generation_error", status="open")
        batch.status = "canceled"
        self.db.commit()
        comments = list(self.word_comments(export_ordered_bilingual_docx(self.db, batch)[0]).values())
        self.assertEqual([row["text"] for row in comments], ["语言与格式：第3句中文理由"])

    def test_comment_export_rejects_unlocatable_range_and_preserves_comment_free_exports(self):
        from docx import Document
        from app.services.docx_comments import SentenceComment, add_sentence_comments
        document = Document()
        document.add_paragraph("A plain paragraph")
        stream = BytesIO()
        document.save(stream)
        raw = stream.getvalue()
        self.assertEqual(add_sentence_comments(raw, []), raw)
        with self.assertRaisesRegex(ValueError, "无法安全定位"):
            add_sentence_comments(raw, [SentenceComment("missing-segment", 0, "无法定位的理由")])


