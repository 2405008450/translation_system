from __future__ import annotations

import asyncio
from io import BytesIO
import json
from pathlib import Path
from tempfile import TemporaryDirectory
from types import SimpleNamespace
import unittest
from unittest.mock import patch
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from sqlalchemy import create_engine
from sqlalchemy.orm import Session

from app.models import DocumentAlignmentPair, Project, Segment, SegmentComment, SegmentQAIssue, User, TranslationReviewReport, TranslationReviewReportItem
from app.services.document_alignment.export import (
    build_export_readiness, export_document_pair_xlsx, export_target_xlsx,
)
from app.services.document_alignment.parser import parse_side
from app.services.document_alignment.segments import materialize_alignment
from app.services.document_alignment.service import (
    _compute_pairs, _orm_units, _store_pairs, create_alignment_batch,
    preview_document_pair, refresh_pair_text, validate_pair_integrity,
)
from app.services.document_alignment.xlsx import (
    MAIN, NS, build_xlsx_pairs, patch_xlsx_text, read_xlsx_content,
)
from tests.test_proofreading import TABLES, _create_sqlite_schema


def workbook_bytes(*, target=False, reverse=False, merge="A1:B1", drawing=True, gaps=False):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "设备规格"
    sheet["A1"] = "电池设备" if target else "Battery equipment"
    if merge:
        sheet.merge_cells(merge)
    sheet["A3"] = CellRichText(TextBlock(InlineFont(b=True, color="FFFF0000"), "电池" if target else "Battery"),
                             TextBlock(InlineFont(i=True), "设备" if target else " equipment"))
    sheet["B3"] = "=SUM(C3:C4)"
    sheet["C3"] = "123" if target else 123
    sheet["C4"] = "SD_HOST_NP01_01"
    sheet["A8"] = "换行\n内容" if target else "Line one\nLine two"
    if gaps:
        sheet["A8"] = None
        sheet["E8"] = "译文新增说明"
    sheet.column_dimensions["A"].width = 35
    sheet.row_dimensions[1].height = 30
    sheet.freeze_panes = "A3"
    other = workbook.create_sheet("历史")
    other["A1"] = "历史记录" if target else "History"
    if reverse:
        workbook.move_sheet(other, offset=-1)
    stream = BytesIO()
    workbook.save(stream)
    workbook.close()
    if not drawing:
        return stream.getvalue()
    # 最小图形夹具含两个带样式的 run，以及不应被回写改动的图片资源。
    xml = '''<xdr:wsDr xmlns:xdr="http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing"
        xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">
      <xdr:twoCellAnchor><xdr:from><xdr:col>3</xdr:col><xdr:colOff>0</xdr:colOff><xdr:row>4</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:from>
       <xdr:to><xdr:col>6</xdr:col><xdr:row>6</xdr:row></xdr:to>
       <xdr:sp><xdr:nvSpPr><xdr:cNvPr id="17" name="说明"/></xdr:nvSpPr><xdr:spPr/>
        <xdr:txBody><a:bodyPr/><a:lstStyle/><a:p>
         <a:r><a:rPr b="1"/><a:t>FIRST</a:t></a:r><a:r><a:rPr i="1"/><a:t>SECOND</a:t></a:r>
        </a:p></xdr:txBody></xdr:sp><xdr:clientData/>
      </xdr:twoCellAnchor></xdr:wsDr>'''.replace("FIRST", "图形" if target else "Drawing").replace("SECOND", "说明" if target else " text")
    output = BytesIO()
    with ZipFile(stream) as archive, ZipFile(output, "w") as result:
        for info in archive.infolist():
            value = archive.read(info.filename)
            if info.filename == ("xl/worksheets/sheet2.xml" if reverse else "xl/worksheets/sheet1.xml"):
                root = ET.fromstring(value)
                ET.SubElement(root, f"{{{MAIN}}}drawing", {"{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id": "rDrawing"})
                value = ET.tostring(root)
            result.writestr(info, value)
        sheet_part = "sheet2.xml" if reverse else "sheet1.xml"
        result.writestr(f"xl/worksheets/_rels/{sheet_part}.rels", '''<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rDrawing" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/drawing" Target="../drawings/drawing1.xml"/></Relationships>''')
        result.writestr("xl/drawings/drawing1.xml", xml)
        result.writestr("xl/media/image1.png", b"unchanged-media-fixture")
    return output.getvalue()


class XlsxPairParsingTests(unittest.TestCase):
    def test_whole_cell_parsing_protects_numbers_formulas_and_merge_children(self):
        content = read_xlsx_content(workbook_bytes())
        self.assertEqual(content.summary["formula_cells"], 1)
        self.assertEqual(content.summary["merged_regions"], 1)
        self.assertEqual(content.summary["drawing_texts"], 1)
        self.assertEqual(len(content.items), 5)
        units = parse_side(workbook_bytes(), "source.xlsx", "sentence")
        self.assertEqual(len(units), 5)
        self.assertIn("Line one\nLine two", [unit.text for unit in units])

    def test_preview_pairs_by_sheet_name_even_when_sheet_order_changes(self):
        result = preview_document_pair(workbook_bytes(), "source.xlsx", workbook_bytes(target=True, reverse=True), "target.xlsx")
        self.assertEqual(result["coordinate_pairs"], 5)
        self.assertEqual(result["unmatched_units"], 0)
        self.assertEqual(result["structure_warnings"], 0)
        self.assertEqual(result["supported_granularities"], ["paragraph"])

    def test_merge_mismatch_requires_review_and_mixed_formats_are_rejected(self):
        result = preview_document_pair(workbook_bytes(), "source.xlsx", workbook_bytes(target=True, merge="A1:C1"), "target.xlsx")
        self.assertEqual(result["structure_warnings"], 1)
        with self.assertRaisesRegex(ValueError, "均为"):
            preview_document_pair(workbook_bytes(), "source.xlsx", b"Text", "target.txt")

    def test_missing_and_added_cells_remain_separate_gaps(self):
        result = preview_document_pair(workbook_bytes(), "source.xlsx", workbook_bytes(target=True, gaps=True), "target.xlsx")
        self.assertEqual(result["coordinate_pairs"], 4)
        self.assertEqual(result["unmatched_units"], 2)

    def test_reordered_sheets_and_locked_pairs_preserve_complete_coverage(self):
        source = read_xlsx_content(workbook_bytes())
        target = read_xlsx_content(workbook_bytes(target=True, reverse=True))
        config = {"xlsx_manifest": {"source": source.items, "target": target.items}}
        pairs = build_xlsx_pairs(source.units(), target.units(), config)
        locked = [(p.src_indices, p.tgt_indices) for p in pairs if p.src_indices in ([0], [4])]
        rerun = build_xlsx_pairs(source.units(), target.units(), config, locked)
        self.assertEqual(sorted(i for p in rerun for i in p.src_indices), list(range(5)))
        self.assertEqual(sorted(i for p in rerun for i in p.tgt_indices), list(range(5)))
        self.assertEqual(sum(p.method == "manual" for p in rerun), 2)

    def test_rewrite_preserves_zip_parts_formula_merges_styles_and_drawing_runs(self):
        original = workbook_bytes(target=True)
        content = read_xlsx_content(original)
        rich = next(item for item in content.items if item.get("cell_ref") == "A3")
        drawing = next(item for item in content.items if item["kind"] == "drawing_text")
        result = patch_xlsx_text(original, [(rich, "电池设备修订"), (drawing, "图形说明修订")])
        with ZipFile(BytesIO(original)) as before, ZipFile(BytesIO(result)) as after:
            self.assertEqual(before.namelist(), after.namelist())
            changed = {name for name in before.namelist() if before.read(name) != after.read(name)}
            self.assertEqual(changed, {"xl/worksheets/sheet1.xml", "xl/drawings/drawing1.xml"})
            oldroot = ET.fromstring(before.read("xl/worksheets/sheet1.xml"))
            newroot = ET.fromstring(after.read("xl/worksheets/sheet1.xml"))
            for ref in ["A1", "B3", "C3", "C4", "A8"]:
                old = next(c for c in oldroot.findall("s:sheetData/s:row/s:c", NS) if c.get("r") == ref)
                new = next(c for c in newroot.findall("s:sheetData/s:row/s:c", NS) if c.get("r") == ref)
                self.assertEqual(ET.tostring(old), ET.tostring(new))
            richcell = next(c for c in newroot.findall("s:sheetData/s:row/s:c", NS) if c.get("r") == "A3")
            self.assertEqual(len(richcell.findall(".//s:rPr", NS)), 2)
            self.assertEqual(ET.tostring(oldroot.find("s:mergeCells", NS)), ET.tostring(newroot.find("s:mergeCells", NS)))
            drawroot = ET.fromstring(after.read("xl/drawings/drawing1.xml"))
            self.assertEqual(len(drawroot.findall(".//a:rPr", NS)), 2)
        reparsed = read_xlsx_content(result)
        self.assertIn("电池设备修订", [item["text"] for item in reparsed.items])
        self.assertIn("图形说明修订", [item["text"] for item in reparsed.items])
        self.assertEqual(patch_xlsx_text(original, []), original)

    def test_shared_string_edit_does_not_change_other_cells_using_same_string(self):
        original = workbook_bytes(target=True, drawing=False)
        output = BytesIO()
        with ZipFile(BytesIO(original)) as before, ZipFile(output, "w") as after:
            for info in before.infolist():
                value = before.read(info.filename)
                if info.filename == "xl/worksheets/sheet1.xml":
                    root = ET.fromstring(value)
                    for cell in root.findall("s:sheetData/s:row/s:c", NS):
                        if cell.get("r") in {"A1", "A3"}:
                            for child in list(cell):
                                cell.remove(child)
                            cell.set("t", "s")
                            ET.SubElement(cell, f"{{{MAIN}}}v").text = "0"
                    value = ET.tostring(root)
                after.writestr(info, value)
            after.writestr("xl/sharedStrings.xml", f'<sst xmlns="{MAIN}"><si><t>共享文字</t></si></sst>')
        original = output.getvalue()
        content = read_xlsx_content(original)
        item = next(i for i in content.items if i.get("cell_ref") == "A1")
        result = patch_xlsx_text(original, [(item, "只改第一格")])
        values = {i.get("cell_ref"): i["text"] for i in read_xlsx_content(result).items if i["sheet_name"] == "设备规格"}
        self.assertEqual(values["A1"], "只改第一格")
        self.assertEqual(values["A3"], "共享文字")
        with ZipFile(BytesIO(original)) as before, ZipFile(BytesIO(result)) as after:
            self.assertEqual(before.read("xl/sharedStrings.xml"), after.read("xl/sharedStrings.xml"))


class XlsxPairWorkflowTests(unittest.TestCase):
    def setUp(self):
        self.storage = TemporaryDirectory()
        self.settings = SimpleNamespace(file_storage_dir=self.storage.name, alignment_llm_full_review_model="test")
        self.storage_patch = patch("app.services.document_storage.get_settings", return_value=self.settings)
        self.alignment_patch = patch("app.services.document_alignment.service.get_settings", return_value=self.settings)
        self.storage_patch.start()
        self.alignment_patch.start()
        self.engine = create_engine("sqlite:///:memory:", future=True)
        with patch("tests.test_proofreading.TABLES", [*TABLES, SegmentComment.__table__, SegmentQAIssue.__table__]):
            _create_sqlite_schema(self.engine)
        self.db = Session(self.engine)
        self.user = User(username="xlsx-admin", hashed_password="x", role="admin")
        self.project = Project(name="复杂表格校对", workflow_template_id="proofread")
        self.db.add_all([self.user, self.project])
        self.db.commit()
        self.target = workbook_bytes(target=True)
        self.batch = create_alignment_batch(
            self.db, project=self.project, current_user=self.user,
            source_bytes=workbook_bytes(), source_filename="source.xlsx",
            target_bytes=self.target, target_filename="target.xlsx",
            source_language="en-US", target_language="zh-CN",
        )
        src, tgt = _orm_units(self.db, self.batch.id, "source"), _orm_units(self.db, self.batch.id, "target")
        with patch("app.services.document_alignment.service.build_semantic_scorer", side_effect=AssertionError("Excel 不应调用向量服务")):
            pairs = asyncio.run(_compute_pairs(self.batch, src, tgt))
        _store_pairs(self.db, self.batch, pairs)
        validate_pair_integrity(self.db, self.batch.id)
        with patch("app.services.document_alignment.segments.attach_project_reference_bases_to_file"):
            self.file = materialize_alignment(self.db, self.batch, current_user=self.user)

    def tearDown(self):
        self.db.close()
        self.engine.dispose()
        self.storage_patch.stop()
        self.alignment_patch.stop()
        self.storage.cleanup()

    def test_report_stage_target_rewrite_and_queue_integration(self):
        config = json.loads(self.batch.config_json)
        self.assertFalse(config["full_review"])
        self.assertEqual(config["alignment_strategy"], "xlsx_coordinate")
        readiness = build_export_readiness(self.db, self.batch)
        self.assertEqual(readiness["available_formats"], ["proofreading_audit_xlsx", "proofreading_docx_ordered"])
        report, _ = export_document_pair_xlsx(self.db, self.batch)
        workbook = load_workbook(BytesIO(report))
        self.assertNotIn("校对译文", [c.value for c in workbook.active[1]])
        self.assertEqual(workbook.active["C2"].value, "设备规格!A1")
        workbook.close()
        self.batch.workflow_stage = "proofreading"
        segment = self.db.query(Segment).filter_by(file_record_id=self.file.id).order_by(Segment.sequence_index).first()
        segment.target_text = "=这是文字而不是公式"
        segment.status = "confirmed"
        self.db.commit()
        readiness = build_export_readiness(self.db, self.batch)
        self.assertEqual(readiness["available_formats"], ["proofreading_audit_xlsx", "proofreading_xlsx_target", "proofreading_docx_ordered"])
        report, _ = export_document_pair_xlsx(self.db, self.batch)
        workbook = load_workbook(BytesIO(report))
        self.assertEqual(workbook.active["G2"].value, "=这是文字而不是公式")
        self.assertEqual(workbook.active["G2"].data_type, "s")
        self.assertEqual(workbook.active["J2"].value, "人工修改")
        workbook.close()
        from app.services.file_export_queue import build_file_record_exported_file
        exported = build_file_record_exported_file(self.db, self.file, "proofreading_xlsx_target")
        result, filename = exported.content, exported.filename
        self.assertTrue(filename.endswith("_校对版.xlsx"))
        self.assertEqual(read_xlsx_content(result).items[0]["text"], "=这是文字而不是公式")
        with ZipFile(BytesIO(self.target)) as original, ZipFile(BytesIO(result)) as exported:
            self.assertEqual(original.read("xl/drawings/drawing1.xml"), exported.read("xl/drawings/drawing1.xml"))
        path = Path(self.storage.name) / "alignment_sources" / f"{self.batch.id}_target.xlsx"
        self.assertEqual(path.read_bytes(), self.target)

    def test_manual_alignment_text_is_included_in_target_rewrite(self):
        pair = self.db.query(DocumentAlignmentPair).filter_by(batch_id=self.batch.id, pair_order=0).one()
        pair.features = json.dumps({"manual_target_text": "对齐阶段修订"}, ensure_ascii=False)
        refresh_pair_text(self.db, pair)
        from app.services.document_alignment.segments import ensure_document_pair_segments_complete
        ensure_document_pair_segments_complete(self.db, self.batch, refresh_existing=True)
        self.batch.workflow_stage = "proofreading"
        self.db.commit()
        result, _ = export_target_xlsx(self.db, self.batch)
        self.assertEqual(read_xlsx_content(result).items[0]["text"], "对齐阶段修订")

    def test_changed_cross_cell_pair_requires_split_before_writeback(self):
        from app.services.document_alignment.service import merge_alignment_pair_range
        from app.services.document_alignment.segments import ensure_document_pair_segments_complete
        pairs = self.db.query(DocumentAlignmentPair).filter_by(batch_id=self.batch.id).order_by(DocumentAlignmentPair.pair_order).limit(2).all()
        merged = merge_alignment_pair_range(self.db, self.batch.id, [p.id for p in pairs])
        ensure_document_pair_segments_complete(self.db, self.batch, refresh_existing=True)
        segment = next(s for s in self.db.query(Segment).filter_by(file_record_id=self.file.id).all()
                       if json.loads(s.segment_metadata)["alignment_pair_id"] == str(merged.id))
        segment.target_text += "人工修改"
        self.batch.workflow_stage = "proofreading"
        self.db.commit()
        with self.assertRaisesRegex(ValueError, "没有唯一"):
            export_target_xlsx(self.db, self.batch)
        # 回写被阻止时，独立报告仍可完整导出。
        content, _ = export_document_pair_xlsx(self.db, self.batch)
        workbook = load_workbook(BytesIO(content))
        self.assertIn("人工修改", workbook.active["G2"].value)
        workbook.close()

    def test_latest_matching_reason_and_workbench_export_format_are_exposed(self):
        self.batch.workflow_stage = "proofreading"
        segment = self.db.query(Segment).filter_by(file_record_id=self.file.id).order_by(Segment.sequence_index).first()
        segment.target_text = "电池生产设备"
        report = TranslationReviewReport(proofreading_batch_id=self.batch.id, project_id=self.project.id)
        self.db.add(report)
        self.db.flush()
        self.db.add(TranslationReviewReportItem(report_id=report.id, file_record_id=self.file.id,
                    segment_id=segment.id, suggested_target_text=segment.target_text, reason="统一设备名称"))
        self.db.commit()
        content, _ = export_document_pair_xlsx(self.db, self.batch)
        workbook = load_workbook(BytesIO(content))
        self.assertEqual(workbook.active["J2"].value, "统一设备名称")
        workbook.close()
        from app.routers.proofreading import get_file_proofreading_baselines
        from app.routers.api import _serialize_file_proofreading_summary
        context = get_file_proofreading_baselines(self.file.id, self.db, self.user)
        self.assertEqual(context["proofreading_context"]["pair_format"], "xlsx")
        self.assertEqual(_serialize_file_proofreading_summary(self.batch)["pair_format"], "xlsx")

    def test_rerun_preserves_locked_manual_text_and_location(self):
        pair = self.db.query(DocumentAlignmentPair).filter_by(batch_id=self.batch.id, pair_order=0).one()
        pair.locked = True
        pair.features = json.dumps({"manual_target_text": "锁定的人工修订"}, ensure_ascii=False)
        refresh_pair_text(self.db, pair)
        self.db.commit()
        config = json.loads(self.batch.config_json)
        source, target = _orm_units(self.db, self.batch.id, "source"), _orm_units(self.db, self.batch.id, "target")
        anchors = [(json.loads(pair.src_indices), json.loads(pair.tgt_indices))]
        pairs = build_xlsx_pairs(source, target, config, anchors)
        signatures = {(tuple(s), tuple(t)) for s, t in anchors}
        _store_pairs(self.db, self.batch, pairs, locked_signatures=signatures)
        validate_pair_integrity(self.db, self.batch.id)
        restored = self.db.query(DocumentAlignmentPair).filter_by(batch_id=self.batch.id, pair_order=0).one()
        self.assertEqual(restored.target_text, "锁定的人工修订")
        self.assertTrue(restored.locked)
        self.assertEqual(json.loads(restored.features)["target_location"], "设备规格!A1")


if __name__ == "__main__":
    unittest.main()
