from __future__ import annotations

import unittest
import json
from io import BytesIO
from tempfile import TemporaryDirectory
from types import SimpleNamespace
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import Session

from app.models import (
    DocumentAlignmentPair,
    DocumentAlignmentUnit,
    FileRecord,
    Project,
    ProjectWorkflowStep,
    ProofreadingBatch,
    ProofreadingColumnBinding,
    ProofreadingSegmentBaseline,
    Segment,
    SegmentRevision,
    TranslationReviewReport,
    TranslationReviewReportItem,
    User,
)
from app.services.proofreading import (
    _pack_groups,
    create_batch_from_workbook,
    export_batch_xlsx,
    generate_batch,
    preview_workbook,
)
from app.services.document_alignment.export import (
    build_export_readiness,
    build_proofreading_export_rows,
    export_document_pair_xlsx,
    export_ordered_bilingual_docx,
)
from app.services.document_alignment.service import (
    delete_alignment_pair,
    merge_alignment_pair_range,
    replace_alignment_pair_range,
    validate_pair_integrity,
)
from app.services.document_alignment.segments import ensure_document_pair_segments_complete
from app.services.file_record_service import get_segment_ordering_for_file_record
from app.routers.proofreading import get_file_proofreading_baselines
from app.routers.api import _apply_segment_scope_filter


TABLES = [
    User.__table__,
    Project.__table__,
    ProjectWorkflowStep.__table__,
    FileRecord.__table__,
    Segment.__table__,
    SegmentRevision.__table__,
    ProofreadingBatch.__table__,
    DocumentAlignmentPair.__table__,
    DocumentAlignmentUnit.__table__,
    ProofreadingColumnBinding.__table__,
    ProofreadingSegmentBaseline.__table__,
    TranslationReviewReport.__table__,
    TranslationReviewReportItem.__table__,
]


def _create_sqlite_schema(engine) -> None:
    # 生产模型的 UUID 默认值是 PostgreSQL 表达式；测试库由 ORM 的 Python 默认值生成 UUID。
    saved_defaults = []
    for table in TABLES:
        for column in table.columns:
            server_default = column.server_default
            if server_default is not None and "::uuid" in str(server_default.arg):
                saved_defaults.append((column, server_default))
                column.server_default = None
    try:
        for table in TABLES:
            table.create(engine)
    finally:
        for column, server_default in saved_defaults:
            column.server_default = server_default


def _workbook_bytes(*, with_formula: bool = False) -> bytes:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "描述"
    worksheet.column_dimensions["B"].width = 28
    worksheet.append(["描述 - 英语 (English)", "描述 - 法语 (French)", "描述 - 德语 (German)"])
    worksheet.append(["Total Distance", "Distance totale", "Gesamtstrecke"])
    worksheet.append(["Total Distance", "Distance total", ""])
    if with_formula:
        worksheet["D2"] = "=LEN(A2)"
    stream = BytesIO()
    workbook.save(stream)
    workbook.close()
    return stream.getvalue()


class ProofreadingPreviewTests(unittest.TestCase):
    def test_preview_detects_columns_languages_and_samples(self):
        result = preview_workbook(_workbook_bytes(), "sample.xlsx")

        self.assertEqual(result["sheets"][0]["header_row"], 1)
        self.assertTrue(result["sheets"][0]["supported"])
        columns = result["sheets"][0]["columns"]
        self.assertEqual([item["suggested_language"] for item in columns], ["en-US", "fr-FR", "de-DE"])
        self.assertEqual(columns[0]["suggested_role"], "source")
        self.assertEqual(columns[1]["samples"][0], "Distance totale")

    def test_preview_blocks_formula_workbooks(self):
        result = preview_workbook(_workbook_bytes(with_formula=True), "complex.xlsx")

        self.assertFalse(result["sheets"][0]["supported"])
        self.assertIn("公式", result["sheets"][0]["blocked_reasons"][0])

    def test_dynamic_batches_have_hard_unique_group_limit(self):
        groups = [
            {"source_text": f"source-{index}", "variants": ["target"], "sid": str(index)}
            for index in range(45)
        ]

        packed = _pack_groups(groups)

        self.assertEqual([len(item) for item in packed], [20, 20, 5])


class ProofreadingWorkflowTests(unittest.IsolatedAsyncioTestCase):
    async def asyncSetUp(self):
        self.storage = TemporaryDirectory()
        self.storage_patch = patch(
            "app.services.document_storage.get_settings",
            return_value=SimpleNamespace(file_storage_dir=self.storage.name),
        )
        self.storage_patch.start()
        self.engine = create_engine("sqlite:///:memory:", future=True)
        _create_sqlite_schema(self.engine)
        self.db = Session(self.engine)
        self.user = User(username="admin", hashed_password="x", role="admin")
        self.project = Project(name="多语校对", workflow_template_id="proofread")
        self.db.add_all([self.user, self.project])
        self.db.flush()
        self.db.add(ProjectWorkflowStep(
            project_id=self.project.id,
            step_key="proofread",
            name="校对",
            step_type="proofread",
            sort_order=0,
        ))
        self.db.commit()

    async def asyncTearDown(self):
        self.db.close()
        self.engine.dispose()
        self.storage_patch.stop()
        self.storage.cleanup()

    async def test_manual_pair_range_merge_is_atomic_locked_and_contiguous(self):
        batch = ProofreadingBatch(
            project_id=self.project.id,
            created_by_id=self.user.id,
            filename="source.docx",
            file_hash="merge-test",
            source_language="zh-CN",
            target_language="en-US",
            batch_kind="document_pair",
            alignment_status="draft",
            status="draft",
        )
        self.db.add(batch)
        self.db.flush()
        source_units = [
            DocumentAlignmentUnit(
                batch_id=batch.id, side="source", unit_index=0, para_index=0,
                block_type="paragraph", block_index=9, text="原文内容",
            ),
        ]
        target_units = [
            DocumentAlignmentUnit(
                batch_id=batch.id, side="target", unit_index=index, para_index=index,
                block_type="paragraph", block_index=index, text=text,
            )
            for index, text in enumerate(["Addition A", "Addition B", "Translation"])
        ]
        pairs = [
            DocumentAlignmentPair(
                batch_id=batch.id, pair_order=0, src_indices="[]", tgt_indices="[0]",
                source_text="", target_text="Addition A", confidence=0.35,
                confidence_level="low", method="dp", features="{}",
            ),
            DocumentAlignmentPair(
                batch_id=batch.id, pair_order=1, src_indices="[]", tgt_indices="[1]",
                source_text="", target_text="Addition B", confidence=0.35,
                confidence_level="low", method="dp", features="{}",
            ),
            DocumentAlignmentPair(
                batch_id=batch.id, pair_order=2, src_indices="[0]", tgt_indices="[2]",
                source_text="原文内容", target_text="Translation", confidence=0.35,
                confidence_level="low", method="dp", features="{}",
            ),
        ]
        self.db.add_all([*source_units, *target_units, *pairs])
        self.db.flush()

        with self.assertRaisesRegex(ValueError, "连续"):
            merge_alignment_pair_range(self.db, batch.id, [pairs[0].id, pairs[2].id])

        merged = merge_alignment_pair_range(self.db, batch.id, [pair.id for pair in pairs])

        remaining = self.db.query(DocumentAlignmentPair).filter_by(batch_id=batch.id).all()
        self.assertEqual(len(remaining), 1)
        self.assertEqual(merged.pair_order, 0)
        self.assertEqual(json.loads(merged.src_indices), [0])
        self.assertEqual(json.loads(merged.tgt_indices), [0, 1, 2])
        self.assertEqual(merged.source_text, "原文内容")
        self.assertEqual(merged.target_text, "Addition A\nAddition B\nTranslation")
        self.assertTrue(merged.locked)
        self.assertEqual(merged.method, "manual")
        self.assertEqual(merged.block_index, 9)

    async def test_delete_translation_only_pair_absorbs_units_without_restoring_text(self):
        batch = ProofreadingBatch(
            project_id=self.project.id,
            created_by_id=self.user.id,
            filename="source.docx",
            file_hash="delete-addition-test",
            source_language="en-US",
            target_language="de-DE",
            batch_kind="document_pair",
            alignment_status="draft",
            workflow_stage="alignment",
            status="draft",
        )
        self.db.add(batch)
        self.db.flush()
        self.db.add_all([
            DocumentAlignmentUnit(
                batch_id=batch.id,
                side="source",
                unit_index=index,
                para_index=index,
                block_type="paragraph",
                block_index=index,
                text=f"Source {index}",
            )
            for index in range(2)
        ] + [
            DocumentAlignmentUnit(
                batch_id=batch.id,
                side="target",
                unit_index=index,
                para_index=index,
                block_type="paragraph",
                block_index=index,
                text=text,
            )
            for index, text in enumerate([
                "Übersetzung A",
                "Fälschlich als Zusatz erkannt",
                "Übersetzung B",
            ])
        ])
        pairs = [
            DocumentAlignmentPair(
                batch_id=batch.id,
                pair_order=0,
                src_indices="[0]",
                tgt_indices="[0]",
                source_text="Source 0",
                target_text="Übersetzung A",
                confidence=1.0,
                confidence_level="high",
                method="manual",
                features="{}",
            ),
            DocumentAlignmentPair(
                batch_id=batch.id,
                pair_order=1,
                src_indices="[]",
                tgt_indices="[1]",
                source_text="",
                target_text="",
                confidence=1.0,
                confidence_level="high",
                method="manual",
                features=json.dumps({"manual_target_text": ""}),
            ),
            DocumentAlignmentPair(
                batch_id=batch.id,
                pair_order=2,
                src_indices="[1]",
                tgt_indices="[2]",
                source_text="Source 1",
                target_text="Übersetzung B",
                confidence=1.0,
                confidence_level="high",
                method="manual",
                features="{}",
            ),
        ]
        self.db.add_all(pairs)
        self.db.flush()

        neighbor = delete_alignment_pair(self.db, batch.id, pairs[1].id)
        self.db.commit()

        remaining = self.db.query(DocumentAlignmentPair).filter_by(batch_id=batch.id).order_by(
            DocumentAlignmentPair.pair_order,
        ).all()
        self.assertEqual(len(remaining), 2)
        self.assertEqual([row.pair_order for row in remaining], [0, 1])
        self.assertEqual(json.loads(remaining[0].tgt_indices), [0, 1])
        self.assertEqual(remaining[0].target_text, "Übersetzung A")
        self.assertEqual(neighbor.id, remaining[0].id)
        self.assertTrue(json.loads(remaining[0].features)["manual_deleted_translation_only"])
        validate_pair_integrity(self.db, batch.id)

    async def test_entering_proofreading_refreshes_only_alignment_changed_segments(self):
        batch = ProofreadingBatch(
            project_id=self.project.id,
            created_by_id=self.user.id,
            filename="source.docx",
            file_hash="alignment-refresh-test",
            source_language="en-US",
            target_language="de-DE",
            batch_kind="document_pair",
            alignment_status="confirmed",
            workflow_stage="alignment",
            status="ready",
        )
        file_record = FileRecord(
            project_id=self.project.id,
            filename="source.docx",
            file_hash="alignment-refresh-test",
            status="in_progress",
            creator_id=self.user.id,
            source_language="en-US",
            target_language="de-DE",
            document_parse_mode="full",
            document_parse_options="{}",
        )
        self.db.add_all([batch, file_record])
        self.db.flush()
        binding = ProofreadingColumnBinding(
            batch_id=batch.id,
            file_record_id=file_record.id,
            sheet_index=0,
            sheet_name="双文档对齐",
            header_row=1,
            source_column=1,
            target_column=2,
            output_column=3,
            source_header="原文",
            target_header="译文",
            target_language="de-DE",
        )
        self.db.add(binding)
        self.db.flush()
        changed_pair = DocumentAlignmentPair(
            batch_id=batch.id,
            pair_order=0,
            src_indices="[0]",
            tgt_indices="[0]",
            source_text="Changed source",
            target_text="Neu ausgerichtete Übersetzung",
            confidence=1.0,
            confidence_level="high",
            method="manual",
            features="{}",
        )
        unchanged_pair = DocumentAlignmentPair(
            batch_id=batch.id,
            pair_order=1,
            src_indices="[1]",
            tgt_indices="[1]",
            source_text="Unchanged source",
            target_text="Unveränderte Basis",
            confidence=1.0,
            confidence_level="high",
            method="manual",
            features="{}",
        )
        self.db.add_all([changed_pair, unchanged_pair])
        self.db.flush()

        segments = []
        for pair, original_target, current_target in (
            (changed_pair, "Alte Ausrichtung", "Veraltete Korrektur"),
            (unchanged_pair, "Unveränderte Basis", "Bewusste manuelle Korrektur"),
        ):
            segment = Segment(
                file_record_id=file_record.id,
                workflow_step_id=self.project.workflow_steps[0].id,
                sentence_id=f"align-{pair.id}",
                source_text=pair.source_text,
                display_text=pair.source_text,
                target_text=current_target,
                status="confirmed",
                source="manual",
                source_word_count=2,
                block_type="paragraph",
                block_index=pair.pair_order,
                sequence_index=pair.pair_order,
                display_index=pair.pair_order,
                segment_metadata=json.dumps({
                    "alignment_pair_id": str(pair.id),
                    "src_indices": json.loads(pair.src_indices),
                    "tgt_indices": json.loads(pair.tgt_indices),
                }),
            )
            self.db.add(segment)
            self.db.flush()
            self.db.add(ProofreadingSegmentBaseline(
                batch_id=batch.id,
                binding_id=binding.id,
                segment_id=segment.id,
                sheet_index=0,
                row_index=pair.pair_order,
                source_cell_ref=f"S{pair.pair_order}",
                target_cell_ref=f"T{pair.pair_order}",
                original_target_text=original_target,
            ))
            self.db.add(SegmentRevision(
                file_record_id=file_record.id,
                segment_id=segment.id,
                sentence_id=segment.sentence_id,
                before_text=original_target,
                after_text=current_target,
                source="manual",
                status="pending",
                author_id=self.user.id,
            ))
            segments.append(segment)
        self.db.commit()

        ensure_document_pair_segments_complete(self.db, batch, refresh_existing=True)
        self.db.commit()

        self.db.refresh(segments[0])
        self.db.refresh(segments[1])
        changed_baseline = self.db.query(ProofreadingSegmentBaseline).filter_by(
            segment_id=segments[0].id,
        ).one()
        self.assertEqual(segments[0].target_text, "Neu ausgerichtete Übersetzung")
        self.assertEqual(changed_baseline.original_target_text, "Neu ausgerichtete Übersetzung")
        self.assertEqual(segments[0].status, "none")
        self.assertEqual(
            self.db.query(SegmentRevision).filter_by(segment_id=segments[0].id).count(),
            0,
        )

        self.assertEqual(segments[1].target_text, "Bewusste manuelle Korrektur")
        self.assertEqual(segments[1].status, "confirmed")
        self.assertEqual(
            self.db.query(SegmentRevision).filter_by(segment_id=segments[1].id).count(),
            1,
        )

    async def test_manual_pair_range_replace_supports_undo_and_empty_rows(self):
        batch = ProofreadingBatch(
            project_id=self.project.id,
            created_by_id=self.user.id,
            filename="source.docx",
            file_hash="replace-test",
            source_language="zh-CN",
            target_language="en-US",
            batch_kind="document_pair",
            alignment_status="draft",
            status="draft",
        )
        self.db.add(batch)
        self.db.flush()
        self.db.add_all([
            DocumentAlignmentUnit(
                batch_id=batch.id, side="source", unit_index=index, para_index=index,
                block_type="paragraph", block_index=index, text=f"原文{index}",
            )
            for index in range(2)
        ] + [
            DocumentAlignmentUnit(
                batch_id=batch.id, side="target", unit_index=index, para_index=index,
                block_type="paragraph", block_index=index, text=f"Target {index}",
            )
            for index in range(2)
        ] + [
            DocumentAlignmentPair(
                batch_id=batch.id, pair_order=index,
                src_indices=json.dumps([index]), tgt_indices=json.dumps([index]),
                source_text=f"原文{index}", target_text=f"Target {index}",
                confidence=0.8, confidence_level="high", method="dp", features="{}",
            )
            for index in range(2)
        ])
        self.db.flush()

        replace_alignment_pair_range(self.db, batch.id, 1, 0, [{
            "src_indices": [], "tgt_indices": [], "locked": True,
        }])
        rows = self.db.query(DocumentAlignmentPair).filter_by(batch_id=batch.id).order_by(
            DocumentAlignmentPair.pair_order,
        ).all()
        self.assertEqual(len(rows), 3)
        self.assertEqual((json.loads(rows[1].src_indices), json.loads(rows[1].tgt_indices)), ([], []))

        # 模拟撤回：用空 replacement 删除刚插入的空行。
        replace_alignment_pair_range(self.db, batch.id, 1, 1, [])
        rows = self.db.query(DocumentAlignmentPair).filter_by(batch_id=batch.id).order_by(
            DocumentAlignmentPair.pair_order,
        ).all()
        self.assertEqual([row.pair_order for row in rows], [0, 1])
        self.assertEqual([json.loads(row.src_indices) for row in rows], [[0], [1]])
        self.assertEqual([json.loads(row.tgt_indices) for row in rows], [[0], [1]])

        # 合并后再按历史快照恢复，覆盖 N→M 与 M→N 两种替换。
        before = [
            {"src_indices": [0], "tgt_indices": [0], "locked": False},
            {"src_indices": [1], "tgt_indices": [1], "locked": False},
        ]
        replace_alignment_pair_range(self.db, batch.id, 0, 2, [{
            "src_indices": [0, 1], "tgt_indices": [0, 1], "locked": True,
        }])
        replace_alignment_pair_range(self.db, batch.id, 0, 1, before)
        rows = self.db.query(DocumentAlignmentPair).filter_by(batch_id=batch.id).order_by(
            DocumentAlignmentPair.pair_order,
        ).all()
        self.assertEqual([json.loads(row.src_indices) for row in rows], [[0], [1]])
        self.assertEqual([row.source_text for row in rows], ["原文0", "原文1"])

    async def test_document_pair_export_manifest_backfills_addition_in_pair_order(self):
        batch = ProofreadingBatch(
            project_id=self.project.id,
            created_by_id=self.user.id,
            filename="source.docx",
            file_hash="test-hash",
            source_language="zh-CN",
            target_language="en-US",
            batch_kind="document_pair",
            alignment_status="confirmed",
            status="completed",
        )
        file_record = FileRecord(
            project_id=self.project.id,
            filename="source.docx",
            file_hash="test-hash",
            status="in_progress",
            creator_id=self.user.id,
            source_language="zh-CN",
            target_language="en-US",
            document_parse_mode="full",
            document_parse_options="{}",
        )
        self.db.add_all([batch, file_record])
        self.db.flush()
        binding = ProofreadingColumnBinding(
            batch_id=batch.id,
            file_record_id=file_record.id,
            sheet_index=0,
            sheet_name="双文档对齐",
            header_row=1,
            source_column=1,
            target_column=2,
            output_column=3,
            source_header="原文",
            target_header="译文",
            target_language="en-US",
        )
        self.db.add(binding)
        self.db.flush()
        pairs = [
            DocumentAlignmentPair(
                batch_id=batch.id, pair_order=0, src_indices="[0]", tgt_indices="[0]",
                source_text="第一条原文。", target_text="First translation.",
                confidence=0.95, confidence_level="high", method="llm_full_review",
                block_type="paragraph", block_index=0,
            ),
            DocumentAlignmentPair(
                batch_id=batch.id, pair_order=1, src_indices="[]", tgt_indices="[1]",
                source_text="", target_text="Translation-only glossary.",
                confidence=0.35, confidence_level="low", method="dp",
                block_type="paragraph", block_index=0,
            ),
            DocumentAlignmentPair(
                batch_id=batch.id, pair_order=2, src_indices="[1]", tgt_indices="[]",
                source_text="第三条原文。", target_text="",
                confidence=0.35, confidence_level="low", method="dp",
                block_type="paragraph", block_index=1,
            ),
        ]
        self.db.add_all(pairs)
        self.db.flush()
        for sequence, pair in ((0, pairs[0]), (1, pairs[2])):
            segment = Segment(
                file_record_id=file_record.id,
                workflow_step_id=self.project.workflow_steps[0].id,
                sentence_id=f"align-{pair.pair_order:05d}",
                source_text=pair.source_text,
                display_text=pair.source_text,
                target_text="Manually reviewed translation." if pair.pair_order == 0 else pair.target_text,
                status="confirmed" if pair.pair_order == 0 else "none",
                source="imported_translation" if pair.target_text else "none",
                source_word_count=1,
                block_type=pair.block_type,
                block_index=pair.block_index,
                sequence_index=sequence,
                display_index=sequence,
                segment_metadata=json.dumps({"alignment_pair_id": str(pair.id)}),
            )
            self.db.add(segment)
            self.db.flush()
            self.db.add(ProofreadingSegmentBaseline(
                batch_id=batch.id,
                binding_id=binding.id,
                segment_id=segment.id,
                sheet_index=0,
                row_index=pair.pair_order,
                source_cell_ref=f"S{pair.pair_order}",
                target_cell_ref=f"T{pair.pair_order}",
                original_target_text=pair.target_text,
            ))
        self.db.commit()

        baseline_payload = get_file_proofreading_baselines(file_record.id, self.db, self.user)
        self.assertTrue(baseline_payload["is_proofreading"])
        self.assertEqual(len(baseline_payload["items"]), 3)
        self.assertEqual(
            self.db.query(Segment).filter_by(file_record_id=file_record.id).count(),
            3,
        )

        rows, _ = build_proofreading_export_rows(self.db, batch)
        self.assertEqual([row.order for row in rows], [0, 1, 2])
        self.assertEqual([row.kind for row in rows], ["对齐", "增译", "缺译"])
        self.assertEqual(rows[0].original_target_text, "First translation.")
        self.assertEqual(rows[0].reviewed_target_text, "Manually reviewed translation.")
        self.assertTrue(rows[0].changed)
        self.assertEqual(rows[1].source_text, "（增译，无对应原文）")
        self.assertEqual(self.db.query(Segment).filter_by(file_record_id=file_record.id).count(), 3)
        self.assertEqual(
            [row.sequence_index for row in self.db.query(Segment).filter_by(
                file_record_id=file_record.id,
            ).order_by(Segment.sequence_index)],
            [0, 1, 2],
        )

        readiness = build_export_readiness(self.db, batch)
        self.assertEqual(readiness["total"], 3)
        self.assertEqual(readiness["translation_only"], 1)
        self.assertEqual(readiness["missing_translation"], 1)
        self.assertTrue(readiness["has_warnings"])

        xlsx_content, _ = export_document_pair_xlsx(self.db, batch)
        workbook = load_workbook(BytesIO(xlsx_content))
        try:
            sheet = workbook["双文档对齐"]
            self.assertEqual(sheet.max_row, 4)
            self.assertEqual(
                [sheet.cell(row=index, column=2).value for index in range(2, 5)],
                ["已对齐", "译文无对应原文", "原文无对应译文"],
            )
            self.assertEqual(sheet.cell(row=2, column=4).value, "First translation.")
            self.assertNotEqual(sheet.cell(row=2, column=4).value, "Manually reviewed translation.")
            self.assertIsNone(sheet.cell(row=4, column=4).value)
        finally:
            workbook.close()

        docx_content, _ = export_ordered_bilingual_docx(self.db, batch)
        from docx import Document
        document = Document(BytesIO(docx_content))
        text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        self.assertLess(text.index("Manually reviewed translation."), text.index("Translation-only glossary."))
        self.assertLess(text.index("Translation-only glossary."), text.index("【译文缺失】"))

    async def test_xlsx_export_manifest_orders_by_sheet_row_then_target_column(self):
        mappings = [{
            "sheet_index": 0,
            "header_row": 1,
            "source_column": 1,
            "targets": [
                {"target_column": 3, "target_language": "de-DE"},
                {"target_column": 2, "target_language": "fr-FR"},
            ],
        }]
        with patch("app.services.proofreading.attach_project_reference_bases_to_file"):
            batch = create_batch_from_workbook(
                self.db,
                project=self.project,
                current_user=self.user,
                raw_bytes=_workbook_bytes(),
                filename="sample.xlsx",
                source_language="en-US",
                mappings=mappings,
            )

        rows, _ = build_proofreading_export_rows(self.db, batch)

        self.assertEqual(
            [(row.row_index, row.cell_index) for row in rows],
            [(2, 1), (2, 2), (3, 1), (3, 2)],
        )

    async def test_create_generate_and_export_multilingual_batch(self):
        mappings = [{
            "sheet_index": 0,
            "header_row": 1,
            "source_column": 1,
            "targets": [
                {"target_column": 2, "target_language": "fr-FR"},
                {"target_column": 3, "target_language": "de-DE"},
            ],
        }]
        with patch("app.services.proofreading.attach_project_reference_bases_to_file"):
            batch = create_batch_from_workbook(
                self.db,
                project=self.project,
                current_user=self.user,
                raw_bytes=_workbook_bytes(),
                filename="sample.xlsx",
                source_language="en-US",
                mappings=mappings,
            )

        self.assertEqual(self.db.query(FileRecord).count(), 2)
        self.assertEqual(self.db.query(ProofreadingColumnBinding).count(), 2)
        self.assertEqual(self.db.query(Segment).count(), 4)
        self.assertEqual(self.db.query(ProofreadingSegmentBaseline).count(), 4)
        self.assertEqual(batch.skipped_segments, 1)

        calls: list[tuple[str, list[dict], str]] = []

        async def fake_review(groups, *, target_language, user_instructions="", **_kwargs):
            calls.append((target_language, groups, user_instructions))
            text_by_language = {"fr-FR": "Distance totale standard", "de-DE": "Standard-Gesamtstrecke"}
            return ({
                group["sid"]: {
                    "reviewed_target_text": text_by_language[target_language],
                    "reason": "统一术语",
                    "category": "consistency",
                    "confidence": "high",
                }
                for group in groups
            }, "test", "test-model")

        with (
            patch("app.services.proofreading._review_group_batch", side_effect=fake_review),
            patch("app.services.proofreading.record_translation_metric_event"),
            patch("app.services.proofreading.sync_file_record_status"),
        ):
            await generate_batch(
                self.db,
                batch,
                current_user=self.user,
                user_instructions="Keep vehicle diagnostic abbreviations unchanged.",
            )

        self.assertEqual(len(calls), 2)
        self.assertEqual(sorted(len(groups) for _, groups, _ in calls), [1, 2])
        self.assertTrue(all(instructions.startswith("Keep vehicle") for _, _, instructions in calls))
        french = self.db.query(FileRecord).filter(FileRecord.target_language == "fr-FR").one()
        french_targets = [row.target_text for row in self.db.query(Segment).filter(
            Segment.file_record_id == french.id,
        ).order_by(Segment.sequence_index)]
        self.assertEqual(french_targets, ["Distance totale standard", "Distance totale standard"])
        self.assertEqual(self.db.query(SegmentRevision).count(), 3)
        self.assertEqual(batch.changed_segments, 3)
        self.assertEqual(batch.status, "completed")

        baseline_payload = get_file_proofreading_baselines(french.id, self.db, self.user)
        self.assertTrue(baseline_payload["is_proofreading"])
        self.assertEqual(
            [item["review_suggestion"] for item in baseline_payload["items"]],
            ["统一术语", "统一术语"],
        )
        self.assertTrue(all(
            item["review_category"] == "consistency"
            for item in baseline_payload["items"]
        ))

        exported, filename = export_batch_xlsx(self.db, batch)
        workbook = load_workbook(BytesIO(exported))
        try:
            worksheet = workbook["描述"]
            self.assertEqual(filename, "sample_校对版.xlsx")
            self.assertEqual(
                [worksheet.cell(row=1, column=index).value for index in range(1, 6)],
                [
                    "描述 - 英语 (English)",
                    "描述 - 法语 (French)",
                    "描述 - 法语 (French)（校对版）",
                    "描述 - 德语 (German)",
                    "描述 - 德语 (German)（校对版）",
                ],
            )
            self.assertEqual(worksheet["B2"].value, "Distance totale")
            self.assertEqual(worksheet["C2"].value, "Distance totale standard")
            self.assertEqual(worksheet["C2"].fill.fgColor.rgb, "FFDDEBFF")
            self.assertEqual(worksheet["C2"].font.color.rgb, "FF0563C1")
            self.assertTrue(worksheet["C2"].font.bold)
            self.assertEqual(worksheet["D2"].value, "Gesamtstrecke")
            self.assertEqual(worksheet["E2"].value, "Standard-Gesamtstrecke")
            self.assertEqual(worksheet["E2"].fill.fgColor.rgb, "FFDDEBFF")
            self.assertIsNone(worksheet["E3"].value)
            self.assertEqual(worksheet.column_dimensions["C"].width, 28)
        finally:
            workbook.close()

    async def test_changed_scope_hides_unchanged_proofreading_segments(self):
        mappings = [{
            "sheet_index": 0,
            "header_row": 1,
            "source_column": 1,
            "targets": [{"target_column": 2, "target_language": "fr-FR"}],
        }]
        with patch("app.services.proofreading.attach_project_reference_bases_to_file"):
            create_batch_from_workbook(
                self.db,
                project=self.project,
                current_user=self.user,
                raw_bytes=_workbook_bytes(),
                filename="sample.xlsx",
                source_language="en-US",
                mappings=mappings,
            )

        french = self.db.query(FileRecord).filter(FileRecord.target_language == "fr-FR").one()
        segments = self.db.query(Segment).filter(
            Segment.file_record_id == french.id,
        ).order_by(Segment.sequence_index).all()
        segments[0].target_text = "Distance totale révisée"
        self.db.commit()

        changed_segments = _apply_segment_scope_filter(
            self.db.query(Segment).filter(Segment.file_record_id == french.id),
            "proofreading_changed",
        ).all()

        self.assertEqual([segment.id for segment in changed_segments], [segments[0].id])

    async def test_document_pair_scopes_filter_translation_only_and_missing_translation(self):
        file_record = FileRecord(
            project_id=self.project.id,
            filename="aligned.docx",
            file_hash="alignment-scope-test",
            status="in_progress",
        )
        self.db.add(file_record)
        self.db.flush()
        aligned = Segment(
            file_record_id=file_record.id,
            sentence_id="align-00000",
            source_text="原文",
            display_text="原文",
            target_text="Translation",
            sequence_index=0,
            display_index=0,
        )
        translation_only = Segment(
            file_record_id=file_record.id,
            sentence_id="align-00001",
            source_text="（增译，无对应原文）",
            display_text="（增译，无对应原文）",
            target_text="Additional translation",
            sequence_index=1,
            display_index=1,
        )
        missing_translation = Segment(
            file_record_id=file_record.id,
            sentence_id="align-00002",
            source_text="漏译原文",
            display_text="漏译原文",
            target_text="",
            sequence_index=2,
            display_index=2,
        )
        self.db.add_all([aligned, translation_only, missing_translation])
        self.db.commit()

        base_query = self.db.query(Segment).filter(Segment.file_record_id == file_record.id)
        translation_only_rows = _apply_segment_scope_filter(
            base_query,
            "proofreading_translation_only",
        ).all()
        missing_translation_rows = _apply_segment_scope_filter(
            base_query,
            "proofreading_missing_translation",
        ).all()

        self.assertEqual([segment.id for segment in translation_only_rows], [translation_only.id])
        self.assertEqual([segment.id for segment in missing_translation_rows], [missing_translation.id])

    async def test_document_pair_workbench_order_ignores_cross_document_blocks(self):
        file_record = FileRecord(
            project_id=self.project.id,
            filename="aligned.docx",
            file_hash="document-pair-ordering",
            status="in_progress",
            document_parse_options=json.dumps({"alignment_mode": "document_pair"}),
        )
        self.db.add(file_record)
        self.db.flush()
        missing = Segment(
            file_record_id=file_record.id,
            sentence_id="align-00000",
            source_text="漏译原文",
            display_text="漏译原文",
            target_text="",
            block_index=900,
            sequence_index=0,
            display_index=0,
        )
        addition = Segment(
            file_record_id=file_record.id,
            sentence_id="align-00001",
            source_text="（增译，无对应原文）",
            display_text="（增译，无对应原文）",
            target_text="Additional translation",
            block_index=1,
            sequence_index=1,
            display_index=1,
        )
        self.db.add_all([missing, addition])
        self.db.commit()

        rows = self.db.query(Segment).filter(
            Segment.file_record_id == file_record.id,
        ).order_by(*get_segment_ordering_for_file_record(file_record)).all()

        self.assertEqual([row.id for row in rows], [missing.id, addition.id])

    async def test_failed_scope_and_retry_only_process_latest_generation_errors(self):
        mappings = [{
            "sheet_index": 0,
            "header_row": 1,
            "source_column": 1,
            "targets": [{"target_column": 2, "target_language": "fr-FR"}],
        }]
        with patch("app.services.proofreading.attach_project_reference_bases_to_file"):
            batch = create_batch_from_workbook(
                self.db,
                project=self.project,
                current_user=self.user,
                raw_bytes=_workbook_bytes(),
                filename="sample.xlsx",
                source_language="en-US",
                mappings=mappings,
            )

        french = self.db.query(FileRecord).filter(FileRecord.target_language == "fr-FR").one()
        segments = self.db.query(Segment).filter(
            Segment.file_record_id == french.id,
        ).order_by(Segment.sequence_index).all()
        failed_segment = segments[1]
        previous_report = TranslationReviewReport(
            project_id=self.project.id,
            created_by_id=self.user.id,
            scope="proofreading_batch",
            report_mode="proofread_generate",
            proofreading_batch_id=batch.id,
            status="partial_failed",
        )
        self.db.add(previous_report)
        self.db.flush()
        self.db.add(TranslationReviewReportItem(
            report_id=previous_report.id,
            project_id=self.project.id,
            file_record_id=french.id,
            segment_id=failed_segment.id,
            sentence_id=failed_segment.sentence_id,
            file_name=french.filename,
            category_key="generation_error",
            reason="模型返回格式不完整",
        ))
        batch.status = "partial_failed"
        batch.failed_segments = 1
        self.db.commit()

        failed_segments = _apply_segment_scope_filter(
            self.db.query(Segment).filter(Segment.file_record_id == french.id),
            "proofreading_failed",
        ).all()
        self.assertEqual([segment.id for segment in failed_segments], [failed_segment.id])

        calls: list[list[dict]] = []

        async def fake_review(groups, **_kwargs):
            calls.append(groups)
            return ({
                group["sid"]: {
                    "reviewed_target_text": group["variants"][0],
                    "reason": "重试成功",
                    "category": "consistency",
                    "confidence": "high",
                }
                for group in groups
            }, "test", "test-model")

        with (
            patch("app.services.proofreading._review_group_batch", side_effect=fake_review),
            patch("app.services.proofreading.sync_file_record_status"),
        ):
            await generate_batch(
                self.db,
                batch,
                current_user=self.user,
                retry_scope="failed_only",
            )

        self.assertEqual(len(calls), 1)
        self.assertEqual(
            [segment.id for group in calls[0] for _, segment in group["segments"]],
            [failed_segment.id],
        )
        self.assertEqual(batch.status, "completed")
        self.assertEqual(batch.failed_segments, 0)
        remaining_failures = _apply_segment_scope_filter(
            self.db.query(Segment).filter(Segment.file_record_id == french.id),
            "proofreading_failed",
        ).all()
        self.assertEqual(remaining_failures, [])

    async def test_failed_only_retry_keeps_earlier_review_suggestions(self):
        mappings = [{
            "sheet_index": 0,
            "header_row": 1,
            "source_column": 1,
            "targets": [{"target_column": 2, "target_language": "fr-FR"}],
        }]
        with patch("app.services.proofreading.attach_project_reference_bases_to_file"):
            batch = create_batch_from_workbook(
                self.db,
                project=self.project,
                current_user=self.user,
                raw_bytes=_workbook_bytes(),
                filename="sample.xlsx",
                source_language="en-US",
                mappings=mappings,
            )
        french = self.db.query(FileRecord).filter(FileRecord.target_language == "fr-FR").one()
        segments = self.db.query(Segment).filter(
            Segment.file_record_id == french.id,
        ).order_by(Segment.sequence_index).all()
        baselines = {
            row.segment_id: row.original_target_text
            for row in self.db.query(ProofreadingSegmentBaseline).filter(
                ProofreadingSegmentBaseline.batch_id == batch.id,
            )
        }

        first_report = TranslationReviewReport(
            project_id=self.project.id,
            created_by_id=self.user.id,
            scope="proofreading_batch",
            report_mode="proofread_generate",
            proofreading_batch_id=batch.id,
            status="partial_failed",
        )
        self.db.add(first_report)
        self.db.flush()
        for segment, category, reason in (
            (segments[0], "terminology", "术语统一"),
            (segments[1], "generation_error", "模型返回格式不完整"),
        ):
            self.db.add(TranslationReviewReportItem(
                report_id=first_report.id,
                project_id=self.project.id,
                file_record_id=french.id,
                segment_id=segment.id,
                sentence_id=segment.sentence_id,
                file_name=french.filename,
                category_key=category,
                reason=reason,
            ))
        segments[0].target_text = "Distance totale standard"
        batch.status = "partial_failed"
        batch.failed_segments = 1
        self.db.commit()

        async def fake_review(groups, **_kwargs):
            return ({
                group["sid"]: {
                    "reviewed_target_text": baselines[group["segments"][0][1].id],
                    "reason": "",
                    "category": "consistency",
                    "confidence": "high",
                }
                for group in groups
            }, "test", "test-model")

        with (
            patch("app.services.proofreading._review_group_batch", side_effect=fake_review),
            patch("app.services.proofreading.sync_file_record_status"),
        ):
            await generate_batch(self.db, batch, current_user=self.user, retry_scope="failed_only")

        payload = get_file_proofreading_baselines(french.id, self.db, self.user)
        by_segment = {item["segment_id"]: item for item in payload["items"]}
        self.assertEqual(by_segment[str(segments[0].id)]["review_suggestion"], "术语统一")
        self.assertEqual(by_segment[str(segments[0].id)]["review_category"], "terminology")
        self.assertEqual(by_segment[str(segments[1].id)]["review_category"], "")
        self.assertEqual(batch.failed_segments, 0)

    async def test_generation_skips_segments_edited_during_llm_call(self):
        mappings = [{
            "sheet_index": 0,
            "header_row": 1,
            "source_column": 1,
            "targets": [{"target_column": 2, "target_language": "fr-FR"}],
        }]
        with patch("app.services.proofreading.attach_project_reference_bases_to_file"):
            batch = create_batch_from_workbook(
                self.db,
                project=self.project,
                current_user=self.user,
                raw_bytes=_workbook_bytes(),
                filename="sample.xlsx",
                source_language="en-US",
                mappings=mappings,
            )
        french = self.db.query(FileRecord).filter(FileRecord.target_language == "fr-FR").one()
        segment = self.db.query(Segment).filter(
            Segment.file_record_id == french.id,
        ).order_by(Segment.sequence_index).first()

        async def fake_review(groups, **_kwargs):
            # 模拟模型调用期间用户在工作台保存了人工修改。
            segment.target_text = "Édition manuelle"
            segment.version = int(segment.version or 1) + 1
            self.db.commit()
            return ({
                group["sid"]: {
                    "reviewed_target_text": "Texte du modèle",
                    "reason": "润色",
                    "category": "fluency",
                    "confidence": "high",
                }
                for group in groups
            }, "test", "test-model")

        with (
            patch("app.services.proofreading._review_group_batch", side_effect=fake_review),
            patch("app.services.proofreading.record_translation_metric_event"),
            patch("app.services.proofreading.sync_file_record_status"),
        ):
            await generate_batch(self.db, batch, current_user=self.user)

        self.db.refresh(segment)
        self.assertEqual(segment.target_text, "Édition manuelle")


def _target_docx_bytes() -> bytes:
    from docx import Document

    document = Document()
    document.add_paragraph(
        "All copies must be signed by the chairman. This agreement takes effect on signing. Third sentence here."
    )
    document.add_paragraph("Second paragraph.")
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "Cell one is here. Cell two is here."
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


class TargetRevisionPayloadTests(unittest.TestCase):
    """目标原格式修订导出：只有真实修改的句段才能出现修订。"""

    UNIT_TEXTS = [
        "All copies must be signed by the chairman.",
        "This agreement takes effect on signing.",
        "Third sentence here.",
        "Second paragraph.",
        "Cell one is here.\nCell two is here.",
    ]
    PAIR_TARGET_INDICES = [[0, 1], [2], [3], [4]]

    def setUp(self):
        self.engine = create_engine("sqlite:///:memory:", future=True)
        _create_sqlite_schema(self.engine)
        self.db = Session(self.engine)
        user = User(username="admin", hashed_password="x", role="admin")
        project = Project(name="修订导出", workflow_template_id="proofread")
        self.db.add_all([user, project])
        self.db.flush()
        self.batch = ProofreadingBatch(
            project_id=project.id, created_by_id=user.id, filename="source.docx",
            file_hash="revision-payload", source_language="zh-CN", target_language="en-US",
            batch_kind="document_pair", alignment_status="confirmed", status="ready",
            config_json=json.dumps({"granularity": "sentence"}),
        )
        self.db.add(self.batch)
        self.db.flush()
        self.pairs = []
        for order, indices in enumerate(self.PAIR_TARGET_INDICES):
            pair = DocumentAlignmentPair(
                batch_id=self.batch.id, pair_order=order, src_indices=json.dumps([order]),
                tgt_indices=json.dumps(indices), source_text=f"原文{order}",
                target_text="\n".join(self.UNIT_TEXTS[index] for index in indices),
                confidence=1.0, confidence_level="high", method="dp", features="{}",
            )
            self.pairs.append(pair)
        self.db.add_all(self.pairs)
        self.db.flush()
        self.target_bytes = _target_docx_bytes()

    def tearDown(self):
        self.db.close()
        self.engine.dispose()

    def _payload(
        self,
        overrides: dict[int, str] | None = None,
        baselines: dict[int, str] | None = None,
    ):
        """overrides 为校对阶段的当前译文；baselines 为对齐结束时的基线译文（默认等于目标原件）。"""
        from app.services.document_alignment.export import _build_target_revision_payload

        overrides = overrides or {}
        baselines = baselines or {}
        rows = []
        for order, pair in enumerate(self.pairs):
            baseline = baselines.get(order, pair.target_text)
            rows.append(SimpleNamespace(
                pair_id=str(pair.id),
                original_target_text=baseline,
                reviewed_target_text=overrides.get(order, baseline),
            ))
        with patch(
            "app.services.document_alignment.export.build_proofreading_export_rows",
            return_value=(rows, None),
        ):
            return _build_target_revision_payload(self.db, self.batch, self.target_bytes, "target.docx")

    def test_split_reviewed_text_keeps_unchanged_multi_unit_pairs_intact(self):
        from app.services.document_alignment.export import _split_reviewed_text

        parts = ["一切副本必须经公司董事长签字确认。", "第二句。"]

        self.assertEqual(_split_reviewed_text(parts, "\n".join(parts)), parts)
        self.assertEqual(_split_reviewed_text(parts, "一切副本须经董事长签字。\n第二句。"), ["一切副本须经董事长签字。", "第二句。"])
        self.assertEqual(_split_reviewed_text(parts[:1], "单句"), ["单句"])

    def test_unchanged_pairs_produce_no_revisions(self):
        segments, revisions = self._payload()

        self.assertEqual(revisions, [])
        self.assertEqual(
            [segment["target_text"] for segment in segments],
            [
                "All copies must be signed by the chairman.",
                "This agreement takes effect on signing.",
                "Third sentence here.",
                "Second paragraph.",
                "Cell one is here.",
                "Cell two is here.",
            ],
        )

    def test_single_changed_sentence_does_not_touch_neighbours(self):
        _, revisions = self._payload({1: "Third sentence changed."})

        self.assertEqual(
            [(item["sentence_id"], item["before_text"], item["after_text"]) for item in revisions],
            [("sent-00003", "Third sentence here.", "Third sentence changed.")],
        )

    def test_changing_second_unit_of_multi_unit_pair_leaves_first_unit_untouched(self):
        _, revisions = self._payload({
            0: "All copies must be signed by the chairman.\nThis agreement takes effect upon signature.",
        })

        self.assertEqual(
            [(item["sentence_id"], item["after_text"]) for item in revisions],
            [("sent-00002", "This agreement takes effect upon signature.")],
        )

    def test_table_cell_edit_only_marks_the_edited_line(self):
        _, revisions = self._payload({3: "Cell one was edited.\nCell two is here."})

        self.assertEqual(
            [(item["sentence_id"], item["after_text"]) for item in revisions],
            [("sent-00005", "Cell one was edited.")],
        )

    def test_failed_generation_segment_with_unchanged_text_has_no_revision(self):
        # 生成失败的句段译文保持基线不变，导出时不能因此被标成修订。
        _, revisions = self._payload({0: self.pairs[0].target_text})

        self.assertEqual(revisions, [])

    def test_alignment_stage_edit_is_written_silently_without_revision(self):
        aligned = "Third sentence adjusted at alignment."
        segments, revisions = self._payload(baselines={1: aligned})

        self.assertEqual(revisions, [])
        self.assertEqual(
            {segment["sentence_id"]: segment["target_text"] for segment in segments}["sent-00003"],
            aligned,
        )

    def test_alignment_edit_followed_by_proofreading_edit_uses_aligned_text_as_before(self):
        segments, revisions = self._payload(
            {1: "Third sentence proofread."},
            baselines={1: "Third sentence adjusted at alignment."},
        )

        self.assertEqual(
            [(item["sentence_id"], item["before_text"], item["after_text"]) for item in revisions],
            [("sent-00003", "Third sentence adjusted at alignment.", "Third sentence proofread.")],
        )
        self.assertEqual(
            {segment["sentence_id"]: segment["target_text"] for segment in segments}["sent-00003"],
            "Third sentence proofread.",
        )

    def test_missing_translation_filled_at_alignment_is_inserted_silently(self):
        extra = DocumentAlignmentPair(
            batch_id=self.batch.id, pair_order=99, src_indices=json.dumps([99]),
            tgt_indices=json.dumps([]), source_text="漏译原文", target_text="Inserted at alignment.",
            confidence=1.0, confidence_level="high", method="manual", features="{}",
        )
        self.db.add(extra)
        self.db.flush()
        self.pairs.append(extra)

        segments, revisions = self._payload()

        self.assertEqual(revisions, [])
        self.assertTrue(any("Inserted at alignment." in segment["target_text"] for segment in segments))

    def test_multi_unit_alignment_edit_on_second_unit_and_proofreading_edit_on_first(self):
        aligned = "All copies must be signed by the chairman.\nThis agreement takes effect upon signature."
        current = "All copies must be signed by the board chairman.\nThis agreement takes effect upon signature."
        segments, revisions = self._payload({0: current}, baselines={0: aligned})

        self.assertEqual(
            [(item["sentence_id"], item["before_text"], item["after_text"]) for item in revisions],
            [(
                "sent-00001",
                "All copies must be signed by the chairman.",
                "All copies must be signed by the board chairman.",
            )],
        )
        self.assertEqual(
            {segment["sentence_id"]: segment["target_text"] for segment in segments}["sent-00002"],
            "This agreement takes effect upon signature.",
        )


class ProofreadingReviewGroupTests(unittest.IsolatedAsyncioTestCase):
    @staticmethod
    def _group(sid: str, source: str, variant: str) -> dict:
        return {
            "sid": sid,
            "source_text": source,
            "variants": [variant],
            "segments": [],
            "translation_only": False,
        }

    async def _run(self, groups, contents):
        from app.services.proofreading import _review_group_batch

        responses = iter(SimpleNamespace(content=content, provider="openrouter", model="m") for content in contents)

        async def fake_completion(**_kwargs):
            return next(responses)

        with patch("app.services.proofreading.request_chat_completion", side_effect=fake_completion) as mocked:
            result = await _review_group_batch(
                groups,
                source_language="zh-CN",
                target_language="en-US",
                rules_text="",
                provider="openrouter",
                model="google/gemini-3.7-flash",
            )
        return result, mocked

    async def test_multiline_source_with_single_line_target_is_accepted(self):
        group = self._group("a", "第一行\n第二行", "First line second line")
        content = json.dumps([{
            "seq": 0, "sid": "a", "reviewed_target_text": "First line, second line",
            "changed": True, "reason": "标点", "category": "grammar", "confidence": "high",
        }], ensure_ascii=False)

        (results, provider, model), _ = await self._run([group], [content])

        self.assertEqual(results["a"]["reviewed_target_text"], "First line, second line")
        self.assertEqual((provider, model), ("openrouter", "m"))

    async def test_line_break_placeholders_round_trip_to_real_line_breaks(self):
        group = self._group("a", "第一行\n第二行", "First line\nSecond line")
        content = json.dumps([{
            "seq": 0, "sid": "a", "reviewed_target_text": "First line⟦LB_1⟧Second line.",
            "changed": True, "reason": "标点", "category": "grammar", "confidence": "high",
        }], ensure_ascii=False)

        (results, _, _), mocked = await self._run([group], [content])

        self.assertEqual(results["a"]["reviewed_target_text"], "First line\nSecond line.")
        prompt = mocked.call_args.kwargs["messages"][1]["content"]
        self.assertIn("First line⟦LB_1⟧Second line", prompt)

    async def test_single_invalid_item_does_not_fail_whole_pack(self):
        good = self._group("a", "你好", "Hello")
        bad = self._group("b", "第一行\n第二行", "First\nSecond")
        first = json.dumps([
            {"seq": 0, "sid": "a", "reviewed_target_text": "Hello!", "changed": True,
             "reason": "标点", "category": "grammar", "confidence": "high"},
            {"seq": 1, "sid": "b", "reviewed_target_text": "First Second", "changed": True,
             "reason": "合并", "category": "format", "confidence": "high"},
        ], ensure_ascii=False)

        (results, _, _), mocked = await self._run([good, bad], [first, first])

        self.assertEqual(results["a"]["reviewed_target_text"], "Hello!")
        self.assertIn("换行", results["b"]["error"])
        self.assertEqual(mocked.call_count, 2)
        retry_prompt = mocked.call_args_list[1].kwargs["messages"][1]["content"]
        self.assertNotIn("你好", retry_prompt)


if __name__ == "__main__":
    unittest.main()
