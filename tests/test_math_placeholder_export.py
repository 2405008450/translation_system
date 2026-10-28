"""公式占位符拼写兼容及原始公式结构保护的导出回归。"""
from io import BytesIO
from xml.etree import ElementTree as ET
from zipfile import ZipFile

import pytest
from docx import Document
from docx.oxml import OxmlElement

from app.services.document_exporter import export_translated_docx
from app.services.document_workspace import parse_docx_workspace


def _document_with_formulas() -> bytes:
    document = Document()
    paragraph = document.add_paragraph("计算")
    for index, value in enumerate(("a", "b")):
        if index:
            paragraph.add_run("加")
        formula = OxmlElement("m:oMath")
        run = OxmlElement("m:r")
        text = OxmlElement("m:t")
        text.text = value
        run.append(text)
        formula.append(run)
        paragraph._p.append(formula)
    paragraph.add_run("的值。")
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def _formulas(raw: bytes) -> list[str]:
    with ZipFile(BytesIO(raw)) as archive:
        root = ET.fromstring(archive.read("word/document.xml"))
    return [ET.tostring(node, encoding="unicode") for node in root.iter(
        "{http://schemas.openxmlformats.org/officeDocument/2006/math}oMath"
    )]


@pytest.mark.parametrize("target", [
    "Calculate ⟦MATH_1⟧ plus ⟦MATH_2⟧.",
    "Calculate [[MATH_1]] plus [[MATH_2]].",
    "Calculate ⟦MATH1⟧ plus ⟦MATH2⟧.",
    "Calculate [[MATH1]] plus [[MATH2]].",
    "Calculate ⟦MATH_1⟧ plus [[MATH2]].",
])
def test_export_restores_original_formulas_for_unambiguous_aliases(target):
    raw = _document_with_formulas()
    source = parse_docx_workspace(raw)["segments"]
    assert len(source) == 1
    exported = export_translated_docx(raw, [dict(source[0], target_text=target)])
    restored = parse_docx_workspace(exported)["segments"]
    assert restored[0]["source_text"] == "Calculate ⟦MATH_1⟧ plus ⟦MATH_2⟧."
    assert _formulas(exported) == _formulas(raw)


@pytest.mark.parametrize("target", [
    "Calculate a plus b.",
    "Calculate ⟦MATH1⟧.",
    "Calculate ⟦MATH2⟧ plus ⟦MATH1⟧.",
    "Calculate ⟦MATH1⟧ plus ⟦MATH1⟧.",
    "Calculate ⟦MATH1⟧ plus ⟦MATH3⟧.",
    "Calculate ⟦MATH1⟧ plus ⟦MATH2⟧ plus ⟦MATH2⟧.",
    "Calculate ⟦MATH1]] plus [[MATH2⟧.",
])
def test_export_still_rejects_ambiguous_or_missing_formula_mapping(target):
    raw = _document_with_formulas()
    source = parse_docx_workspace(raw)["segments"]
    with pytest.raises(ValueError, match=f"数学公式占位符.*{source[0]['sentence_id']}"):
        export_translated_docx(raw, [dict(source[0], target_text=target)])


def _revision_text(raw: bytes, accept: bool) -> str:
    with ZipFile(BytesIO(raw)) as archive:
        root = ET.fromstring(archive.read("word/document.xml"))
    word = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    math = "{http://schemas.openxmlformats.org/officeDocument/2006/math}"
    count = 0

    def collect(node):
        nonlocal count
        if node.tag == word + ("del" if accept else "ins"):
            return ""
        if node.tag == math + "oMath":
            count += 1
            return f"⟦MATH_{count}⟧"
        if node.tag in {word + "t", word + "delText"}:
            return node.text or ""
        return "".join(collect(child) for child in node)

    return collect(root)


@pytest.mark.parametrize("before,after", [
    ("Compute ⟦MATH_1⟧ plus ⟦MATH_2⟧.", "Calculate ⟦MATH_1⟧ and ⟦MATH_2⟧!"),
    ("⟦MATH_1⟧⟦MATH_2⟧", "Before ⟦MATH_1⟧ between ⟦MATH_2⟧ after"),
    ("Before ⟦MATH_1⟧ between ⟦MATH_2⟧ after", "⟦MATH_1⟧⟦MATH_2⟧"),
    ("Calculate ⟦MATH_1⟧ plus ⟦MATH_2⟧.", "Calculate ⟦MATH_1⟧ plus ⟦MATH_2⟧!"),
])
def test_formula_revision_export_accept_and_reject_preserve_text_and_formulas(before, after):
    raw = _document_with_formulas()
    source = parse_docx_workspace(raw)["segments"][0]
    revision = dict(id="formula-review", sentence_id=source["sentence_id"], before_text=before,
                    after_text=after, status="pending", source="manual", author="校对人员",
                    comment_text="统一公式前后术语。")
    exported = export_translated_docx(
        raw, [dict(source, target_text=after)], revisions=[revision],
        include_revision_marks=True, require_revision_marks=True,
    )
    assert _revision_text(exported, True) == after
    assert _revision_text(exported, False) == before
    assert _formulas(exported) == _formulas(raw)
    with ZipFile(BytesIO(exported)) as archive:
        document_xml = archive.read("word/document.xml")
        assert b"MATH_" not in document_xml
        root = ET.fromstring(document_xml)
    word = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    paragraph = next(root.iter(word + "p"))
    children = list(paragraph)
    start = next(root.iter(word + "commentRangeStart"))
    end = next(root.iter(word + "commentRangeEnd"))
    formula_indices = [index for index, node in enumerate(children) if node.tag.endswith("}oMath")]
    assert children.index(start) < min(formula_indices) <= max(formula_indices) < children.index(end)
    revision_ids = [node.get(word + "id") for node in root.iter() if node.tag in {word + "ins", word + "del"}]
    assert len(revision_ids) == len(set(revision_ids))


def test_formula_revision_export_rejects_baseline_with_different_formula_sequence():
    raw = _document_with_formulas()
    source = parse_docx_workspace(raw)["segments"][0]
    after = "Calculate ⟦MATH_1⟧ plus ⟦MATH_2⟧."
    revision = dict(id="bad-formula-review", sentence_id=source["sentence_id"],
                    before_text="Calculate ⟦MATH_2⟧ plus ⟦MATH_1⟧.", after_text=after,
                    status="pending", source="manual", author="校对人员")
    with pytest.raises(ValueError, match="无法安全写入真实校对修订"):
        export_translated_docx(raw, [dict(source, target_text=after)], revisions=[revision],
                               include_revision_marks=True, require_revision_marks=True)


@pytest.mark.parametrize("story_kind", ["header", "footer"])
def test_header_footer_revision_is_retained_without_unsupported_word_comment(story_kind):
    document = Document()
    document.add_paragraph("Body text.")
    getattr(document.sections[0], story_kind).paragraphs[0].text = "Source text."
    stream = BytesIO()
    document.save(stream)
    raw = stream.getvalue()
    source = parse_docx_workspace(raw)["segments"]
    selected = next(segment for segment in source if segment["block_type"] == story_kind)
    segments = [dict(segment, target_text="New translation." if segment is selected else segment["source_text"])
                for segment in source]
    revision = dict(id="story-review", sentence_id=selected["sentence_id"],
                    before_text="Old translation.", after_text="New translation.",
                    status="pending", source="manual", author="校对人员", comment_text="统一术语。")
    exported = export_translated_docx(raw, segments, revisions=[revision],
                                       include_revision_marks=True, require_revision_marks=True)
    with ZipFile(BytesIO(exported)) as archive:
        name = next(name for name in archive.namelist() if name.startswith(f"word/{story_kind}") and name.endswith(".xml"))
        root = ET.fromstring(archive.read(name))
    word = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    assert list(root.iter(word + "ins"))
    assert list(root.iter(word + "del"))
    assert not list(root.iter(word + "commentRangeStart"))
    accepted = "".join(node.text or "" for node in root.iter(word + "t"))
    assert accepted == "New translation."
