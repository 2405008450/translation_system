"""在原 XLSX 包中追加校对列，图片、合并区域及其他模板部件原样保留。"""
from __future__ import annotations

from copy import deepcopy
from io import BytesIO
from typing import Any
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from openpyxl.utils import get_column_letter

from app.services.adapters.xlsx_exporter import _register_namespaces, _sanitize_xml_text, _serialize_xml
from app.services.document_alignment.xlsx import MAIN, NS, REL, _relationships


def append_proofreading_columns(raw_bytes: bytes, workbook: Any, columns: list[dict[str, Any]]) -> bytes:
    """只修改选中工作表的单元格 XML，不用 openpyxl 重新保存模板。"""
    _register_namespaces()
    modified: dict[str, bytes] = {}
    with ZipFile(BytesIO(raw_bytes)) as archive:
        sheets = ET.fromstring(archive.read("xl/workbook.xml")).findall("s:sheets/s:sheet", NS)
        relationships = _relationships(archive, "xl/workbook.xml")
        grouped: dict[int, list[dict[str, Any]]] = {}
        for column in columns:
            grouped.setdefault(column["sheet_index"], []).append(column)
        for sheet_index, entries in grouped.items():
            sheet = workbook.worksheets[sheet_index]
            part = relationships[sheets[sheet_index].get(f"{{{REL}}}id", "")]
            original = archive.read(part)
            root = ET.fromstring(original)
            data = root.find("s:sheetData", NS)
            if data is None:
                raise ValueError(f"工作表“{sheet.title}”缺少单元格数据。")
            rows = {int(row.get("r", "0")): row for row in data}
            source_cells = {cell.get("r"): cell for row in data for cell in row if cell.tag == f"{{{MAIN}}}c"}
            cols = root.find("s:cols", NS)
            if cols is None:
                cols = ET.Element(f"{{{MAIN}}}cols")
                root.insert(list(root).index(data), cols)
            last_column = int(sheet.max_column or 1)
            for entry in sorted(entries, key=lambda item: item["target_column"]):
                last_column += 1
                if last_column > 16384:
                    raise ValueError(f"工作表“{sheet.title}”没有空间追加校对列。")
                letter = get_column_letter(last_column)
                source_column = entry["target_column"]
                # 复制列宽等显示属性，不移动原有列或图片锚点。
                source_dimension = next((col for col in cols if
                    int(col.get("min", "0")) <= source_column <= int(col.get("max", "0"))), None)
                dimension = deepcopy(source_dimension) if source_dimension is not None else ET.Element(f"{{{MAIN}}}col")
                dimension.set("min", str(last_column))
                dimension.set("max", str(last_column))
                if source_dimension is None:
                    dimension.set("width", "28")
                    dimension.set("customWidth", "1")
                cols.append(dimension)
                values = {entry["header_row"]: entry["header"], **entry["values"]}
                for row_index, text in sorted(values.items()):
                    row = rows.get(row_index)
                    if row is None:
                        row = ET.SubElement(data, f"{{{MAIN}}}row", r=str(row_index))
                        rows[row_index] = row
                    cell = ET.Element(f"{{{MAIN}}}c", r=f"{letter}{row_index}", t="inlineStr")
                    source_cell = source_cells.get(f"{get_column_letter(source_column)}{row_index}")
                    if source_cell is not None and source_cell.get("s") is not None:
                        cell.set("s", source_cell.get("s"))
                    container = ET.SubElement(cell, f"{{{MAIN}}}is")
                    node = ET.SubElement(container, f"{{{MAIN}}}t")
                    node.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
                    node.text = _sanitize_xml_text(text)
                    # extLst 必须保留在行末。
                    extension = row.find("s:extLst", NS)
                    row.insert(list(row).index(extension) if extension is not None else len(row), cell)
                    row.attrib.pop("spans", None)
            data[:] = sorted(data, key=lambda row: int(row.get("r", "0")))
            dimension = root.find("s:dimension", NS)
            if dimension is not None:
                dimension.set("ref", f"A1:{get_column_letter(last_column)}{sheet.max_row or 1}")
            modified[part] = _serialize_xml(root, original)
        output = BytesIO()
        with ZipFile(output, "w") as target:
            for info in archive.infolist():
                target.writestr(info, modified.get(info.filename, archive.read(info.filename)))
        return output.getvalue()
