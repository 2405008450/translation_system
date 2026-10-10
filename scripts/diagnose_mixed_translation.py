"""用真实 XLSX 样本对照翻译场景，并检查单元格回写和公式保留。"""
from __future__ import annotations

import argparse
import asyncio
import json
import re
import sys
from dataclasses import replace
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
from zipfile import ZipFile

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from app.services.adapters.xlsx_adapter import XlsxAdapter
from app.services.adapters.xlsx_exporter import XlsxExporter
from app.services.llm_service import LLMTranslationTask, LLMTranslationFailure, iter_batch_translate


DEFAULT_SAMPLES = [
    ("1. System Configuration", "C3"),
    ("1. System Configuration", "B99"),
    ("1. System Configuration", "B100"),
    ("1. System Configuration", "B101"),
    ("2. State & Remote Command", "B216"),
    ("3. Sequence & Scenario List", "B6"),
    ("4. Sequence 별 Data List", "L9"),
]


async def main(args):
    raw = Path(args.xlsx).read_bytes()
    output_dir = Path(args.output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    options = {"xlsx_translate_formula_cells": False}
    parsed = XlsxAdapter().parse_with_options(raw, options=options)
    samples = set(DEFAULT_SAMPLES)
    selected = [segment for segment in parsed.segments
                if (segment.metadata.get("sheet_name"), segment.metadata.get("cell_ref")) in samples]
    assert {(segment.metadata["sheet_name"], segment.metadata["cell_ref"]) for segment in selected} == samples
    tasks = [LLMTranslationTask(
        sentence_id=segment.segment_id, status="none", source_text=segment.source_text,
        source_layout_text=segment.source_layout_text or segment.display_text,
        source_language="ko-KR", target_language="zh-CN", block_type="table_cell",
        block_index=segment.metadata["sheet_index"], row_index=segment.metadata["row"],
        cell_index=segment.metadata["col"],
    ) for segment in selected]
    report = {"samples": len(samples), "segments": len(tasks), "scenarios": {}}
    original = load_workbook(BytesIO(raw), read_only=True, data_only=False)
    try:
        for scenario in ("standard", "mixed"):
            cache_path = output_dir / f"{scenario}.translations.json"
            if scenario == "standard" and cache_path.exists():
                results = [SimpleNamespace(**item) for item in json.loads(cache_path.read_text(encoding="utf-8"))]
                assert {item.sentence_id for item in results} == {task.sentence_id for task in tasks}
                print("复用已有标准场景对照结果", flush=True)
            else:
                print(f"开始 {scenario} 场景，{len(tasks)} 个句段", flush=True)
                results = [item async for item in iter_batch_translate(
                    [replace(task, translation_scenario=scenario) for task in tasks],
                    provider=args.provider, model_override=args.model, translation_unit="sentence",
                )]
            failures = [item for item in results if isinstance(item, LLMTranslationFailure)]
            if failures:
                report["scenarios"][scenario] = {"failures": [item.error_message for item in failures]}
                continue
            by_id = {item.sentence_id: item for item in results}
            (output_dir / f"{scenario}.translations.json").write_text(json.dumps([
                {"sentence_id": item.sentence_id, "translated_text": item.translated_text,
                 "provider": item.provider, "model": item.model} for item in results
            ], ensure_ascii=False, indent=2), encoding="utf-8")
            exported = XlsxExporter().export(raw, [SimpleNamespace(
                sentence_id=segment.segment_id, source_text=segment.source_text,
                target_text=by_id[segment.segment_id].translated_text,
            ) for segment in selected], document_parse_options=options)
            actual = load_workbook(BytesIO(exported), read_only=True, data_only=False)
            try:
                unchanged_errors, formula_count, cells = [], 0, []
                for sheet in original:
                    for row_index, (source_row, target_row) in enumerate(zip(sheet, actual[sheet.title]), 1):
                        for col_index, (source, target) in enumerate(zip(source_row, target_row), 1):
                            coordinate = f"{get_column_letter(col_index)}{row_index}"
                            if source.data_type == "f":
                                formula_count += 1
                            if (sheet.title, coordinate) not in samples and source.value != target.value:
                                unchanged_errors.append(f"{sheet.title}!{coordinate}")
                for sheet_name, cell_ref in DEFAULT_SAMPLES:
                    source = original[sheet_name][cell_ref].value
                    target = actual[sheet_name][cell_ref].value
                    identifiers = re.findall(r"\b[A-Za-z][A-Za-z0-9]*(?:_[A-Za-z0-9]+)+\b", source)
                    cells.append({
                        "sheet": sheet_name, "cell": cell_ref, "source": source, "target": target,
                        "identifiers_preserved": all(target.count(token) == source.count(token) for token in identifiers),
                        "line_breaks_preserved": source.count("\n") == target.count("\n"),
                        "remaining_korean": bool(re.search(r"[\uac00-\ud7af]", target)),
                    })
                with ZipFile(BytesIO(raw)) as before, ZipFile(BytesIO(exported)) as after:
                    untouched_parts = [name for name in before.namelist() if not name.startswith("xl/worksheets/sheet")]
                    package_preserved = all(before.read(name) == after.read(name) for name in untouched_parts)
                report["scenarios"][scenario] = {
                    "model": results[0].model, "provider": results[0].provider,
                    "formula_count": formula_count, "unselected_cell_changes": unchanged_errors,
                    "other_package_parts_preserved": package_preserved, "cells": cells,
                }
            finally:
                actual.close()
            (output_dir / f"{scenario}.xlsx").write_bytes(exported)
            (output_dir / "comparison.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
            print(f"完成 {scenario} 场景", flush=True)
    finally:
        original.close()
    report_path = output_dir / "comparison.json"
    report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"对照报告：{report_path}", flush=True)


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    sys.stderr.reconfigure(encoding="utf-8")
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--xlsx", required=True)
    parser.add_argument("--output-dir", required=True)
    parser.add_argument("--provider", choices=("auto", "deepseek", "openrouter"), default="auto")
    parser.add_argument("--model", default=None)
    asyncio.run(main(parser.parse_args()))
