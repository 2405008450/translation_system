"""脱敏回归：分页边界、句段位置及导出完整性。"""
from io import BytesIO
import re
from xml.etree import ElementTree as ET
from zipfile import ZipFile

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

import app.services.document_exporter as exporter
from app.services.document_workspace import parse_docx_workspace


def source(location="body", kind="rendered", shared=False):
    doc = Document()
    if location == "table":
        paragraph = doc.add_table(rows=1, cols=1).cell(0, 0).paragraphs[0]
    elif location in {"header", "footer"}:
        paragraph = getattr(doc.sections[0], location).paragraphs[0]
    else:
        paragraph = doc.add_paragraph()
    run = paragraph.add_run("首句。分页前半")
    run.bold = True
    for i in range(2):
        marker = OxmlElement("w:lastRenderedPageBreak" if kind == "rendered" else "w:br")
        if kind == "page":
            marker.set(qn("w:type"), "page")
        if shared:
            run._r.append(marker)
            text = OxmlElement("w:t")
            text.text = "分页后半。中句。末段开头" if i == 0 else "末段结尾。最后一句。"
            run._r.append(text)
        else:
            paragraph.add_run()._r.append(marker)
            run = paragraph.add_run("分页后半。中句。末段开头" if i == 0 else "末段结尾。最后一句。")
            run.bold = True
    doc.add_paragraph("下一段落。")
    output = BytesIO()
    doc.save(output)
    return output.getvalue()


def roots(raw):
    with ZipFile(BytesIO(raw)) as archive:
        return {name: ET.fromstring(archive.read(name)) for name in archive.namelist()
                if name.startswith("word/") and name.endswith(".xml")}


def export(raw, rows, mode, options=None):
    if mode == "translated":
        return exporter.export_translated_docx(raw, rows, document_parse_options=options)
    return exporter.export_bilingual_docx_with_layout(raw, rows, order=mode, document_parse_options=options)


@pytest.mark.parametrize("location", ["body", "table", "header", "footer"])
@pytest.mark.parametrize("kind", ["rendered", "page"])
@pytest.mark.parametrize("shared", [False, True])
@pytest.mark.parametrize("mode", ["translated", "source_first", "target_first"])
def test_page_boundary_targets_appear_once_in_order_in_their_story(location, kind, shared, mode):
    raw = source(location, kind, shared)
    rows = [{**row, "sequence_index": i, "target_text": f"TARGET_{i:03d}.", "target_html": f"<b>TARGET_{i:03d}.</b>"}
            for i, row in enumerate(parse_docx_workspace(raw)["segments"])]
    # 页中半句当前按解析器规则合并；导出必须尊重原有句段，不能自行多拆。
    assert len(rows) == 6
    before = roots(raw)
    after = roots(export(raw, rows, mode))
    for name, root in before.items():
        original_paragraphs = [p for p in root.iter(qn("w:p")) if any(t.text for t in p.iter(qn("w:t")))]
        paragraphs = [p for p in after[name].iter(qn("w:p")) if any(t.text for t in p.iter(qn("w:t")))]
        targets = re.findall(r"TARGET_\d{3}", "".join(n.text or "" for n in after[name].iter(qn("w:t"))))
        expected = [f"TARGET_{i:03d}" for i, row in enumerate(rows)
                    if (location not in {"header", "footer"} and name == "word/document.xml")
                    or (location in {"header", "footer"} and
                        ((row["block_type"].startswith(location) and location in name)
                         or (not row["block_type"].startswith(location) and name == "word/document.xml")))]
        assert targets == expected
        factor = 1 if mode == "translated" else 2
        assert len(paragraphs) == factor * len(original_paragraphs)
        for tag in ("lastRenderedPageBreak", "br"):
            assert len(list(after[name].iter(qn("w:" + tag)))) == factor * len(list(root.iter(qn("w:" + tag))))


@pytest.mark.parametrize("mode", ["translated", "source_first", "target_first"])
def test_unvisited_segment_blocks_fail_instead_of_returning_a_partial_file(mode):
    raw = source()
    rows = parse_docx_workspace(raw)["segments"]
    rows.append({"sentence_id": "orphan", "source_text": "不可定位的内容。", "target_text": "Orphan.",
                 "block_type": "paragraph", "block_index": 9999})
    with pytest.raises(ValueError, match="未完整写入"):
        export(raw, rows, mode)


@pytest.mark.parametrize("mode", ["translated", "source_first", "target_first"])
def test_missing_text_token_fails_closed(mode, monkeypatch):
    raw = source()
    rows = parse_docx_workspace(raw)["segments"]
    for row in rows:
        row["target_text"] = "Translated."
    original = exporter._collect_inline_tokens
    def drop_source_text(*args, **kwargs):
        node = kwargs.get("node", args[0] if args else None)
        if node.tag == qn("w:t") and (node.text or "").startswith("首句"):
            return []
        return original(*args, **kwargs)
    monkeypatch.setattr(exporter, "_collect_inline_tokens", drop_source_text)
    with pytest.raises(ValueError, match="句界与译文数量不一致"):
        export(raw, rows, mode)


@pytest.mark.parametrize("mode", ["translated", "source_first", "target_first"])
def test_english_profile_ignores_rendered_breaks_but_preserves_explicit_page_breaks(mode):
    for kind in ("rendered", "page"):
        raw = source(kind=kind)
        options = {"segmentation_profile": "en_docx_v1"}
        rows = parse_docx_workspace(raw, document_parse_options=options)["segments"]
        for i, row in enumerate(rows):
            row["target_text"] = f"TARGET_{i:03d}."
        output = export(raw, rows, mode, options)
        text = "".join(t.text or "" for root in roots(output).values() for t in root.iter(qn("w:t")))
        assert re.findall(r"TARGET_\d{3}", text) == [f"TARGET_{i:03d}" for i in range(len(rows))]


@pytest.mark.parametrize("mode", ["translated", "source_first", "target_first"])
@pytest.mark.parametrize("kind", ["rendered", "page"])
@pytest.mark.parametrize("formatting", ["plain", "rich", "mixed"])
def test_page_marker_that_creates_a_sentence_boundary_keeps_the_final_target(mode, kind, formatting):
    doc = Document()
    paragraph = doc.add_paragraph()
    paragraph.add_run("Section heading")
    marker = OxmlElement("w:lastRenderedPageBreak" if kind == "rendered" else "w:br")
    if kind == "page":
        marker.set(qn("w:type"), "page")
    paragraph.add_run()._r.append(marker)
    paragraph.add_run("This is a much longer explanatory body containing enough words to form a separate sentence. Final sentence.")
    buffer = BytesIO()
    doc.save(buffer)
    raw = buffer.getvalue()
    rows = parse_docx_workspace(raw)["segments"]
    assert len(rows) == 3
    for i, row in enumerate(rows):
        row["sequence_index"] = i
        row["target_text"] = f"TARGET_{i:03d}."
        if formatting == "rich" or (formatting == "mixed" and i % 2):
            row["target_html"] = f"<b>TARGET_{i:03d}.</b>"
    output = export(raw, rows, mode)
    text = "".join(t.text or "" for root in roots(output).values() for t in root.iter(qn("w:t")))
    assert re.findall(r"TARGET_\d{3}", text) == [f"TARGET_{i:03d}" for i in range(3)]


@pytest.mark.parametrize("mode", ["translated", "source_first", "target_first"])
def test_persisted_source_ranges_survive_new_sentence_splitter_rules(mode, monkeypatch):
    raw = source(shared=True)
    rows = parse_docx_workspace(raw)["segments"]
    for i, row in enumerate(rows):
        row["target_text"] = f"TARGET_{i:03d}."
    monkeypatch.setattr(exporter, "split_sentence_spans", lambda text: [exporter.SentenceSpan(0, len(text))])
    output = export(raw, rows, mode)
    text = "".join(t.text or "" for root in roots(output).values() for t in root.iter(qn("w:t")))
    assert re.findall(r"TARGET_\d{3}", text) == [f"TARGET_{i:03d}" for i in range(len(rows))]


def test_revision_markers_in_one_run_preserve_order_and_page_nodes():
    raw = source(shared=True)
    rows = parse_docx_workspace(raw)["segments"]
    revisions = []
    for i, row in enumerate(rows):
        row["target_text"] = f"TARGET_{i:03d}."
        revisions.append({"sentence_id": row["sentence_id"], "id": str(i),
                          "before_text": f"BEFORE_{i:03d}.", "after_text": row["target_text"]})
    output = exporter.export_translated_docx(raw, rows, revisions=revisions, include_revision_marks=True)
    root = roots(output)["word/document.xml"]
    text = "".join(t.text or "" for t in root.iter(qn("w:t")))
    assert re.findall(r"TARGET_\d{3}", text) == [f"TARGET_{i:03d}" for i in range(len(rows))]
    assert len(list(root.iter(qn("w:lastRenderedPageBreak")))) == 2


def test_repeated_writer_visit_is_rejected(monkeypatch):
    raw = source()
    rows = parse_docx_workspace(raw)["segments"]
    original = exporter._mark_segments_consumed
    def twice(segments):
        original(segments)
        original(segments)
    monkeypatch.setattr(exporter, "_mark_segments_consumed", twice)
    with pytest.raises(ValueError, match="重复写入"):
        export(raw, rows, "translated")


@pytest.mark.parametrize("html", ["<b></b>", "<b>Stale incomplete target.</b>"])
def test_stale_rich_text_falls_back_to_the_complete_saved_translation(html):
    raw = source(shared=True)
    rows = parse_docx_workspace(raw)["segments"]
    for i, row in enumerate(rows):
        row["target_text"] = f"TARGET_{i:03d}."
        row["target_html"] = html
    output = export(raw, rows, "translated")
    text = "".join(t.text or "" for root in roots(output).values() for t in root.iter(qn("w:t")))
    assert re.findall(r"TARGET_\d{3}", text) == [f"TARGET_{i:03d}" for i in range(len(rows))]
    assert "Stale incomplete target" not in text


@pytest.mark.parametrize("mode", ["translated", "source_first", "target_first"])
@pytest.mark.parametrize("multiline_indexes", [[3], [1, 3]])
@pytest.mark.parametrize("shared_text_node", [True, False])
def test_later_multiline_targets_keep_their_character_positions(mode, multiline_indexes, shared_text_node):
    doc = Document()
    paragraph = doc.add_paragraph()
    if shared_text_node:
        paragraph.add_run("Alpha sentence. Beta sentence. Gamma sentence. Delta sentence.")
    else:
        run = paragraph.add_run("Alpha sentence. Beta sentence. ")
        text = OxmlElement("w:t")
        text.text = "Gamma sentence. Delta sentence."
        run._r.append(text)
    buffer = BytesIO()
    doc.save(buffer)
    raw = buffer.getvalue()
    rows = parse_docx_workspace(raw)["segments"]
    assert len(rows) == 4
    for i, row in enumerate(rows):
        row["target_text"] = f"TARGET_{i:03d}" + ("\ncontinued." if i in multiline_indexes else ".")
    output = export(raw, rows, mode)
    root = roots(output)["word/document.xml"]
    text = "".join(t.text or "" for t in root.iter(qn("w:t")))
    assert re.findall(r"TARGET_\d{3}", text) == [f"TARGET_{i:03d}" for i in range(4)]
    # 同时验证换行在对应句段内，不能只是正确文字出现在别处。
    stream = "".join((node.text or "") if node.tag == qn("w:t") else "\n"
                     for node in root.iter() if node.tag in {qn("w:t"), qn("w:br")})
    for i in multiline_indexes:
        assert f"TARGET_{i:03d}\ncontinued." in stream


@pytest.mark.parametrize("mode", ["translated", "source_first", "target_first"])
def test_soft_break_in_a_later_sentence_does_not_move_it_before_the_previous_sentence(mode):
    doc = Document()
    paragraph = doc.add_paragraph()
    run = paragraph.add_run("Alpha sentence. Beta sentence with enough words")
    run._r.append(OxmlElement("w:br"))
    tail = OxmlElement("w:t")
    tail.text = " continued. Gamma sentence."
    run._r.append(tail)
    buffer = BytesIO()
    doc.save(buffer)
    raw = buffer.getvalue()
    rows = parse_docx_workspace(raw)["segments"]
    assert len(rows) == 3
    for i, row in enumerate(rows):
        row["target_text"] = f"TARGET_{i:03d}."
    output = export(raw, rows, mode)
    text = "".join(t.text or "" for root in roots(output).values() for t in root.iter(qn("w:t")))
    assert re.findall(r"TARGET_\d{3}", text) == [f"TARGET_{i:03d}" for i in range(3)]


def _swap_final_text(stories, left="TARGET_001", right="TARGET_002"):
    for story in stories:
        for node in story.root.iter(qn("w:t")):
            node.text = (node.text or "").replace(left, "SWAP_TEMP").replace(right, left).replace("SWAP_TEMP", right)


@pytest.mark.parametrize("location", ["body", "table", "header", "footer"])
@pytest.mark.parametrize("mode", ["translated", "source_first", "target_first"])
def test_order_guard_rejects_reordered_content_after_all_writes_and_formatting(location, mode, monkeypatch):
    raw = source(location=location, shared=True)
    rows = parse_docx_workspace(raw)["segments"]
    for i, row in enumerate(rows):
        row["target_text"] = f"TARGET_{i:03d}."
    original = exporter._clean_story_formatting
    def corrupt(stories):
        original(stories)
        _swap_final_text(stories)
        # 所有句段仍存在；只有顺序错误，数量检查本身不能发现。
        text = "".join(node.text or "" for story in stories for node in story.root.iter(qn("w:t")))
        assert sorted(re.findall(r"TARGET_\d{3}", text)) == [f"TARGET_{i:03d}" for i in range(len(rows))]
    monkeypatch.setattr(exporter, "_clean_story_formatting", corrupt)
    with pytest.raises(ValueError, match="句段顺序或位置核验未通过"):
        export(raw, rows, mode, {"clean_format": True})


@pytest.mark.parametrize("mode", ["translated", "source_first", "target_first"])
def test_order_guard_rejects_whole_paragraphs_moved_with_all_their_text(mode, monkeypatch):
    doc = Document()
    doc.add_paragraph("Alpha sentence.")
    doc.add_paragraph("Beta sentence.")
    buffer = BytesIO()
    doc.save(buffer)
    raw = buffer.getvalue()
    rows = parse_docx_workspace(raw)["segments"]
    for i, row in enumerate(rows):
        row["target_text"] = f"TARGET_{i:03d}."
    def corrupt(stories):
        for story in stories:
            for parent in story.root.iter():
                targets = [node for node in parent if node.tag == qn("w:p") and
                           "TARGET_" in "".join(t.text or "" for t in node.iter(qn("w:t")))]
                if len(targets) == 2:
                    left, right = targets
                    first, second = list(parent).index(left), list(parent).index(right)
                    parent.remove(left)
                    parent.remove(right)
                    parent.insert(first, right)
                    parent.insert(second, left)
                    return
        raise AssertionError("No target paragraphs to swap")
    monkeypatch.setattr(exporter, "_clean_story_formatting", corrupt)
    with pytest.raises(ValueError, match="文本块顺序核验未通过"):
        export(raw, rows, mode, {"clean_format": True})


@pytest.mark.parametrize("mode", ["translated", "source_first", "target_first"])
def test_order_guard_rejects_targets_exchanged_between_cells(mode, monkeypatch):
    doc = Document()
    cells = doc.add_table(rows=1, cols=2).rows[0].cells
    cells[0].text, cells[1].text = "Alpha sentence.", "Beta sentence."
    buffer = BytesIO()
    doc.save(buffer)
    raw = buffer.getvalue()
    rows = parse_docx_workspace(raw)["segments"]
    for i, row in enumerate(rows):
        row["target_text"] = f"TARGET_{i:03d}."
    monkeypatch.setattr(exporter, "_clean_story_formatting", lambda stories: _swap_final_text(stories, "TARGET_000", "TARGET_001"))
    with pytest.raises(ValueError, match="句段顺序或位置核验未通过"):
        export(raw, rows, mode, {"clean_format": True})


@pytest.mark.parametrize("mode", ["translated", "source_first", "target_first"])
@pytest.mark.parametrize("missing_repeat", [False, True])
def test_repeated_targets_require_distinct_ordered_occurrences(mode, missing_repeat, monkeypatch):
    doc = Document()
    doc.add_paragraph("Alpha sentence. Beta sentence. Gamma sentence.")
    buffer = BytesIO()
    doc.save(buffer)
    raw = buffer.getvalue()
    rows = parse_docx_workspace(raw)["segments"]
    for row, target in zip(rows, ["Repeat.", "Repeat.", "Tail."], strict=True):
        row["target_text"] = target
    if missing_repeat:
        def remove_one(stories):
            for story in stories:
                for node in story.root.iter(qn("w:t")):
                    if "Repeat." in (node.text or ""):
                        node.text = node.text.replace("Repeat.", "", 1)
                        return
        monkeypatch.setattr(exporter, "_clean_story_formatting", remove_one)
        with pytest.raises(ValueError, match="句段顺序或位置核验未通过"):
            export(raw, rows, mode, {"clean_format": True})
    else:
        export(raw, rows, mode)


@pytest.mark.parametrize("corrupt", [False, True])
def test_revision_deleted_text_cannot_satisfy_the_output_order_check(corrupt, monkeypatch):
    doc = Document()
    doc.add_paragraph("Alpha sentence. Beta sentence.")
    buffer = BytesIO()
    doc.save(buffer)
    raw = buffer.getvalue()
    rows = parse_docx_workspace(raw)["segments"]
    revisions = []
    for i, row in enumerate(rows):
        row["target_text"] = f"TARGET_{i:03d}."
        revisions.append({"sentence_id": row["sentence_id"], "id": str(i),
                          "before_text": f"TARGET_{1-i:03d}.", "after_text": row["target_text"]})
    if corrupt:
        monkeypatch.setattr(exporter, "_clean_story_formatting", lambda stories: _swap_final_text(stories, "TARGET_000", "TARGET_001"))
        with pytest.raises(ValueError, match="句段顺序或位置核验未通过"):
            exporter.export_translated_docx(raw, rows, revisions=revisions, include_revision_marks=True, document_parse_options={"clean_format": True})
    else:
        exporter.export_translated_docx(raw, rows, revisions=revisions, include_revision_marks=True)


@pytest.mark.parametrize("mode", ["translated", "source_first", "target_first"])
def test_order_guard_blocks_the_previous_multiline_insertion_bug(mode, monkeypatch):
    doc = Document()
    doc.add_paragraph("Alpha sentence. Beta sentence. Gamma sentence.")
    buffer = BytesIO()
    doc.save(buffer)
    raw = buffer.getvalue()
    rows = parse_docx_workspace(raw)["segments"]
    for i, row in enumerate(rows):
        row["target_text"] = f"TARGET_{i:03d}" + ("\ncontinued." if i == 2 else ".")
    def legacy_insertion(overlaps, replacement, breaks):
        token = overlaps[0][0]
        parent, anchor = token.container_element, token.anchor_element
        if parent is None or anchor is None:
            return False
        index = list(parent).index(anchor)
        for run in exporter._build_inserted_word_runs(replacement, token.run_element):
            parent.insert(index, run)
            index += 1
        for item, start, end in overlaps:
            item.edits.append((start, end, ""))
        exporter._remove_structural_line_break_tokens(breaks)
        return True
    monkeypatch.setattr(exporter, "_queue_structural_word_text_range_edit", legacy_insertion)
    with pytest.raises(ValueError, match="句段顺序或位置核验未通过") as error:
        export(raw, rows, mode)
    assert "TARGET_" not in str(error.value)
    assert "Alpha" not in str(error.value)


@pytest.mark.parametrize("mode", ["translated", "source_first", "target_first"])
@pytest.mark.parametrize("blank", ["", "   "])
def test_order_guard_accepts_intentional_blank_targets(mode, blank):
    raw = source(shared=True)
    rows = parse_docx_workspace(raw)["segments"]
    for i, row in enumerate(rows):
        row["target_text"] = blank if i % 2 else f"TARGET_{i:03d}."
    export(raw, rows, mode)


@pytest.mark.parametrize("corrupt", [False, True])
def test_order_guard_tracks_real_formula_positions_instead_of_skipping_math(corrupt, monkeypatch):
    doc = Document()
    paragraph = doc.add_paragraph()
    for i, word in enumerate(("Alpha ", "Beta ")):
        paragraph.add_run(word)
        formula = OxmlElement("m:oMath")
        math_run = OxmlElement("m:r")
        math_text = OxmlElement("m:t")
        math_text.text = "x" if i == 0 else "y"
        math_run.append(math_text)
        formula.append(math_run)
        paragraph._p.append(formula)
        paragraph.add_run(". ")
    buffer = BytesIO()
    doc.save(buffer)
    raw = buffer.getvalue()
    rows = parse_docx_workspace(raw)["segments"]
    assert len(rows) == 2
    for i, row in enumerate(rows):
        row["target_text"] = f"TARGET_{i:03d} ⟦MATH_{i+1}⟧."
    if corrupt:
        def swap_formulas(stories):
            for story in stories:
                for parent in story.root.iter():
                    formulas = [child for child in parent if child.tag == qn("m:oMath")]
                    if len(formulas) == 2:
                        a, b = formulas
                        first, second = list(parent).index(a), list(parent).index(b)
                        parent.remove(a)
                        parent.remove(b)
                        parent.insert(first, b)
                        parent.insert(second, a)
                        return
            raise AssertionError("Formula scope not found")
        monkeypatch.setattr(exporter, "_clean_story_formatting", swap_formulas)
        with pytest.raises(ValueError, match="句段顺序或位置核验未通过"):
            export(raw, rows, "translated", {"clean_format": True})
    else:
        export(raw, rows, "translated")


@pytest.mark.parametrize("corrupt", [False, True])
def test_table_line_distribution_has_a_final_order_check(corrupt, monkeypatch):
    doc = Document()
    cell = doc.add_table(rows=1, cols=1).cell(0, 0)
    cell.paragraphs[0].add_run("实施采")
    cell.add_paragraph("购")
    buffer = BytesIO()
    doc.save(buffer)
    raw = buffer.getvalue()
    rows = parse_docx_workspace(raw)["segments"]
    assert len(rows) == 1
    rows[0]["target_text"] = "FIRST_LINE\nSECOND_LINE\nEXTRA_LINE"
    if corrupt:
        def reverse_lines(stories):
            for story in stories:
                for parent in story.root.iter(qn("w:tc")):
                    paragraphs = [p for p in parent if p.tag == qn("w:p")]
                    for p in paragraphs:
                        parent.remove(p)
                    for p in reversed(paragraphs):
                        parent.append(p)
        monkeypatch.setattr(exporter, "_clean_story_formatting", reverse_lines)
        with pytest.raises(ValueError, match="句段顺序或位置核验未通过"):
            export(raw, rows, "translated", {"clean_format": True})
    else:
        export(raw, rows, "translated")


def test_detached_output_scope_is_rejected_even_after_consumption_was_recorded(monkeypatch):
    raw = source()
    rows = parse_docx_workspace(raw)["segments"]
    for i, row in enumerate(rows):
        row["target_text"] = f"TARGET_{i:03d}."
    def detach(stories):
        for story in stories:
            for parent in story.root.iter():
                for child in list(parent):
                    if child.tag == qn("w:p") and "TARGET_000" in "".join(n.text or "" for n in child.iter(qn("w:t"))):
                        parent.remove(child)
                        return
    monkeypatch.setattr(exporter, "_clean_story_formatting", detach)
    with pytest.raises(ValueError, match="输出范围已脱离文档"):
        export(raw, rows, "translated", {"clean_format": True})


@pytest.mark.parametrize("mode", ["translated", "source_first", "target_first"])
def test_full_scope_readback_rejects_duplicate_text_even_when_ordered_search_would_pass(mode, monkeypatch):
    doc = Document()
    doc.add_paragraph("Alpha sentence. Beta sentence.")
    buffer = BytesIO()
    doc.save(buffer)
    raw = buffer.getvalue()
    rows = parse_docx_workspace(raw)["segments"]
    for i, row in enumerate(rows):
        row["target_text"] = f"TARGET_{i:03d}."
    def duplicate(stories):
        for story in stories:
            for node in story.root.iter(qn("w:t")):
                if "TARGET_001." in (node.text or ""):
                    node.text = node.text.replace("TARGET_001.", "TARGET_001.TARGET_001.")
                    return
    monkeypatch.setattr(exporter, "_clean_story_formatting", duplicate)
    with pytest.raises(ValueError, match="输出块与预期内容不一致"):
        export(raw, rows, mode, {"clean_format": True})


def test_order_guard_checks_a_whole_table_against_neighboring_paragraphs(monkeypatch):
    doc = Document()
    doc.add_paragraph("Before the table.")
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "Inside the table."
    doc.add_paragraph("After the table.")
    buffer = BytesIO()
    doc.save(buffer)
    raw = buffer.getvalue()
    rows = parse_docx_workspace(raw)["segments"]
    for i, row in enumerate(rows):
        row["target_text"] = f"TARGET_{i:03d}."
    def move_table(stories):
        for story in stories:
            body = story.root if story.root.tag == qn("w:body") else story.root.find("w:body", exporter.NS)
            if body is not None:
                table = body.find("w:tbl", exporter.NS)
                body.remove(table)
                body.insert(0, table)
    monkeypatch.setattr(exporter, "_clean_story_formatting", move_table)
    with pytest.raises(ValueError, match="段落与表格位置不一致"):
        export(raw, rows, "translated", {"clean_format": True})


@pytest.mark.parametrize("mode", ["translated", "source_first", "target_first"])
@pytest.mark.parametrize("kind", ["footnote", "endnote", "comment"])
def test_multiple_note_stories_in_the_same_part_are_all_checked(mode, kind):
    ET.register_namespace("w", exporter.NS["w"])
    ET.register_namespace("ct", "http://schemas.openxmlformats.org/package/2006/content-types")
    doc = Document()
    doc.add_paragraph("Body sentence.")
    buffer = BytesIO()
    doc.save(buffer)
    payload = BytesIO()
    part = f"word/{kind}s.xml"
    note_root = ET.Element(qn(f"w:{kind}s"))
    for i in (1, 2):
        note = ET.SubElement(note_root, qn(f"w:{kind}"), {qn("w:id"): str(i)})
        paragraph = ET.SubElement(note, qn("w:p"))
        text = ET.SubElement(ET.SubElement(paragraph, qn("w:r")), qn("w:t"))
        text.text = "First note." if i == 1 else "Second note."
    with ZipFile(BytesIO(buffer.getvalue())) as original, ZipFile(payload, "w") as output:
        for name in original.namelist():
            data = original.read(name)
            if name == "[Content_Types].xml":
                root = ET.fromstring(data)
                ET.SubElement(root, "{http://schemas.openxmlformats.org/package/2006/content-types}Override", {
                    "PartName": "/" + part,
                    "ContentType": "application/vnd.openxmlformats-officedocument.wordprocessingml." + kind + "s+xml",
                })
                data = ET.tostring(root)
            output.writestr(name, data)
        output.writestr(part, ET.tostring(note_root))
    raw = payload.getvalue()
    rows = parse_docx_workspace(raw)["segments"]
    assert len(rows) == 3
    for i, row in enumerate(rows):
        row["target_text"] = f"TARGET_{i:03d}."
    result = export(raw, rows, mode)
    content = "".join(n.text or "" for n in roots(result)[part].iter(qn("w:t")))
    assert "TARGET_001" in content and "TARGET_002" in content


@pytest.mark.parametrize("deleted_index", [0, 1])
def test_order_guard_accepts_a_revision_that_intentionally_deletes_the_whole_target(deleted_index):
    doc = Document()
    doc.add_paragraph("Alpha sentence. Beta sentence.")
    buffer = BytesIO()
    doc.save(buffer)
    raw = buffer.getvalue()
    rows = parse_docx_workspace(raw)["segments"]
    for i, row in enumerate(rows):
        row["target_text"] = "" if i == deleted_index else row["source_text"]
    revisions = [{"id": "delete", "sentence_id": rows[deleted_index]["sentence_id"],
                  "before_text": "Baseline target.", "after_text": ""}]
    result = exporter.export_translated_docx(raw, rows, revisions=revisions, include_revision_marks=True)
    text = "".join(n.text or "" for n in roots(result)["word/document.xml"].iter(qn("w:t")))
    assert rows[1-deleted_index]["source_text"] in text
    assert rows[deleted_index]["source_text"] not in text
