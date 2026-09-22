from __future__ import annotations

from io import BytesIO
from xml.etree import ElementTree as ET
from zipfile import ZipFile

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from app.services.document_exporter import (
    NS,
    TextToken,
    _apply_token_edits,
    _queue_text_range_edit,
    export_bilingual_docx_with_layout,
    export_translated_docx,
)
from app.services.document_workspace import parse_docx_workspace


def _source(*, bold=False, location="body", text="Source sentence.", inherited=False):
    document = Document()
    if location == "table":
        paragraph = document.add_table(rows=1, cols=1).cell(0, 0).paragraphs[0]
    elif location == "header":
        paragraph = document.sections[0].header.paragraphs[0]
    elif location == "footer":
        paragraph = document.sections[0].footer.paragraphs[0]
    else:
        paragraph = document.add_paragraph()
    run = paragraph.add_run(text)
    if inherited:
        document.styles["Normal"].font.name = "Arial"
        document.styles["Normal"].font.size = Pt(16)
    else:
        run.font.name = "Arial"
        run.font.size = Pt(16)
        run.font.color.rgb = RGBColor.from_string("C02030")
        fonts = run._r.rPr.rFonts
        fonts.set(qn("w:eastAsia"), "SimSun")
        fonts.set(qn("w:cs"), "Tahoma")
        # 同时保留主题引用及复杂文字字号，防止只恢复西文字体。
        fonts.set(qn("w:asciiTheme"), "minorHAnsi")
        size_cs = OxmlElement("w:szCs")
        size_cs.set(qn("w:val"), "28")
        run._r.rPr.append(size_cs)
    if bold:
        run.bold = True
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _root(raw, part="word/document.xml"):
    with ZipFile(BytesIO(raw)) as archive:
        return ET.fromstring(archive.read(part))


def _runs(raw, part="word/document.xml"):
    return [
        run for run in _root(raw, part).findall(".//w:r", NS)
        if any(node.text for node in run.findall("w:t", NS))
    ]


def _segments(raw, target="Translated sentence.", html=None):
    return [
        {**segment, "target_text": target, "target_html": html}
        for segment in parse_docx_workspace(raw)["segments"]
    ]


def _assert_native_style(source_run, target_run):
    for name in ("rFonts", "sz", "szCs", "color"):
        source = source_run.find(f"w:rPr/w:{name}", NS)
        target = target_run.find(f"w:rPr/w:{name}", NS)
        assert target is not None, name
        assert target.attrib == source.attrib, name


@pytest.mark.parametrize("location", ["body", "table", "header", "footer"])
@pytest.mark.parametrize("bold,html", [(False, None), (True, None), (True, "<i>Translated sentence.</i>")])
def test_stable_export_preserves_native_font_slots_size_and_color(location, bold, html):
    raw = _source(bold=bold, location=location)
    part = {"header": "word/header1.xml", "footer": "word/footer1.xml"}.get(location, "word/document.xml")
    exported = export_translated_docx(raw, _segments(raw, html=html))
    translated = _runs(exported, part)
    assert "".join(t.text or "" for r in translated for t in r.findall("w:t", NS)) == "Translated sentence."
    for run in translated:
        _assert_native_style(_runs(raw, part)[0], run)
        if html:
            assert run.find("w:rPr/w:i", NS) is not None
            assert run.find("w:rPr/w:b", NS) is None


def test_inherited_font_and_size_are_not_materialized_as_export_defaults():
    raw = _source(inherited=True)
    exported = export_translated_docx(raw, _segments(raw))
    for run in _runs(exported):
        assert run.find("w:rPr/w:rFonts", NS) is None
        assert run.find("w:rPr/w:sz", NS) is None
    with ZipFile(BytesIO(raw)) as source, ZipFile(BytesIO(exported)) as target:
        for part in ("word/styles.xml", "word/theme/theme1.xml", "word/fontTable.xml"):
            assert source.read(part) == target.read(part)


def test_explicit_target_font_and_size_override_source():
    raw = _source()
    html = '<span style="font-family:Calibri;font-size:11pt;color:#123456">Translated sentence.</span>'
    exported = export_translated_docx(raw, _segments(raw, html=html))
    run = _runs(exported)[0]
    fonts = run.find("w:rPr/w:rFonts", NS)
    assert set(fonts.attrib.values()) == {"Calibri"}
    assert run.find("w:rPr/w:sz", NS).get(qn("w:val")) == "22"
    assert run.find("w:rPr/w:color", NS).get(qn("w:val")) == "123456"


def test_multiline_insertion_preserves_native_style():
    raw = _source()
    exported = export_translated_docx(raw, _segments(raw, target="First line\nSecond line"))
    assert _root(exported).find(".//w:br", NS) is not None
    for run in _runs(exported):
        _assert_native_style(_runs(raw)[0], run)


def test_consecutive_formatted_sentences_preserve_native_style():
    raw = _source(bold=True, text="First sentence. Second sentence.")
    segments = _segments(raw)
    assert len(segments) == 2
    exported = export_translated_docx(raw, segments)
    for run in _runs(exported):
        _assert_native_style(_runs(raw)[0], run)


def test_clean_format_still_removes_font_and_size_when_requested():
    raw = _source(bold=True)
    exported = export_translated_docx(raw, _segments(raw), document_parse_options={"clean_format": True})
    for run in _runs(exported):
        assert run.find("w:rPr", NS) is None


def test_drawingml_text_replacement_keeps_typefaces_and_theme():
    run = ET.fromstring(f'<a:r xmlns:a="{NS["a"]}"><a:rPr sz="1600"><a:latin typeface="+mn-lt"/><a:ea typeface="SimSun"/><a:cs typeface="Tahoma"/></a:rPr><a:t>Source</a:t></a:r>')
    original = ET.tostring(run.find("a:rPr", NS))
    token = TextToken(element=run.find("a:t", NS), original_text="Source", display_text="Source", source_text="Source", run_element=run)
    _queue_text_range_edit([(token, 0, 6)], "Target")
    _apply_token_edits([token])
    assert run.find("a:t", NS).text == "Target"
    assert ET.tostring(run.find("a:rPr", NS)) == original


def test_bilingual_export_keeps_source_and_target_fonts():
    raw = _source(bold=True)
    exported = export_bilingual_docx_with_layout(raw, _segments(raw))
    text = "".join(t.text or "" for t in _root(exported).findall(".//w:t", NS))
    assert "Source sentence." in text and "Translated sentence." in text
    for run in _runs(exported):
        _assert_native_style(_runs(raw)[0], run)


def test_revision_runs_keep_source_font_and_size():
    raw = _source()
    segments = _segments(raw)
    revision = {"id": "font-test", "sentence_id": segments[0]["sentence_id"], "status": "pending", "source": "manual", "before_text": "Previous sentence.", "after_text": "Translated sentence.", "author": "Test"}
    exported = export_translated_docx(raw, segments, revisions=[revision], include_revision_marks=True)
    assert _root(exported).find(".//w:ins", NS) is not None
    for run in _root(exported).findall(".//w:r", NS):
        if run.find("w:t", NS) is not None or run.find("w:delText", NS) is not None:
            _assert_native_style(_runs(raw)[0], run)


@pytest.mark.parametrize("prefix", ["", "Earlier sentence. "])
def test_mixed_source_anchor_keeps_own_style_without_spreading_it(prefix):
    document = Document()
    paragraph = document.add_paragraph()
    if prefix:
        paragraph.add_run(prefix)
    special = paragraph.add_run("SKU-123")
    special.bold = True
    special.font.name = "Courier New"
    special.font.size = Pt(20)
    special.font.color.rgb = RGBColor.from_string("FF0000")
    normal = paragraph.add_run(" is the product identifier for this item.")
    normal.font.name = "Arial"
    normal.font.size = Pt(12)
    output = BytesIO()
    document.save(output)
    raw = output.getvalue()
    segments = parse_docx_workspace(raw)["segments"]
    target = "The product identifier is SKU-123 for this translated item."
    for segment in segments:
        segment["target_text"] = target if "SKU-123" in segment["source_text"] else segment["source_text"]
    exported = export_translated_docx(raw, segments)
    for run in _runs(exported):
        text = "".join(t.text or "" for t in run.findall("w:t", NS))
        if text == prefix or text == prefix.strip():
            continue
        fonts = run.find("w:rPr/w:rFonts", NS)
        if text == "SKU-123":
            assert fonts.get(qn("w:ascii")) == "Courier New"
            assert run.find("w:rPr/w:sz", NS).get(qn("w:val")) == "40"
        else:
            assert fonts.get(qn("w:ascii")) == "Arial"
            assert run.find("w:rPr/w:sz", NS).get(qn("w:val")) == "24"
            assert run.find("w:rPr/w:color", NS) is None


def test_numbering_localization_preserves_native_number_font():
    document = Document()
    document.add_paragraph("Source sentence.", style="List Number")
    for level in document.part.numbering_part.element.findall(".//" + qn("w:lvl")):
        properties = level.find(qn("w:rPr"))
        if properties is None:
            properties = OxmlElement("w:rPr")
            level.append(properties)
        fonts = properties.find(qn("w:rFonts"))
        if fonts is None:
            fonts = OxmlElement("w:rFonts")
            properties.append(fonts)
        fonts.set(qn("w:ascii"), "Courier New")
        fonts.set(qn("w:asciiTheme"), "minorHAnsi")
        level.find(qn("w:lvlText")).set(qn("w:val"), "第%1章")
    output = BytesIO()
    document.save(output)
    raw = output.getvalue()
    exported = export_translated_docx(raw, _segments(raw), target_language="en-US")
    source = _root(raw, "word/numbering.xml")
    target = _root(exported, "word/numbering.xml")
    assert [n.attrib for n in source.findall(".//w:rFonts", NS)] == [n.attrib for n in target.findall(".//w:rFonts", NS)]
    assert [n.attrib for n in source.findall(".//w:lvlText", NS)] != [n.attrib for n in target.findall(".//w:lvlText", NS)]


def test_enabled_export_table_style_can_still_override_font_and_size():
    from app.services.export_settings.style_export_integration import apply_export_style_settings

    raw = _source(location="table", bold=True)
    exported = export_translated_docx(raw, _segments(raw))
    configured = apply_export_style_settings(exported, {
        "enabled": True,
        "styles": {"Table Grid": {"tbl_run_font_ascii": "Calibri", "tbl_run_font_size": 11}},
    })
    run = _runs(configured)[0]
    assert run.find("w:rPr/w:rFonts", NS).get(qn("w:ascii")) == "Calibri"
    assert run.find("w:rPr/w:sz", NS).get(qn("w:val")) == "22"
